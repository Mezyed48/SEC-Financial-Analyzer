"""
SEC 10-K Filing Downloader & Financial Analyzer - Steps 2 through 9 (Complete Project)

This comprehensive accounting workflow automates:
- Steps 2 & 3: Ticker input validation and SEC EDGAR Form 10-K retrieval.
- Step 4: Management's Discussion & Analysis (Item 7 MD&A) plain-text extraction.
- Step 5: Primary Financial Statements extraction from Item 8 (Income Statement,
  Balance Sheet, and Statement of Cash Flows) into a professionally formatted
  Excel workbook ('output/<TICKER>_Financial_Analysis.xlsx').
- Step 6: Text cleaning and normalization of extracted MD&A text files
  ('cleaned_text/<TICKER>/<TICKER>_<YEAR>_MDA_CLEAN.txt') preparing them for
  Loughran-McDonald financial sentiment analysis.
- Step 7: Loughran-McDonald financial sentiment analysis on cleaned MD&A files:
  * Dynamically detects and loads the Loughran-McDonald dictionary (.csv or .xlsx)
    from the 'dictionaries/' folder.
  * Identifies available sentiment categories (Positive, Negative, Uncertainty,
    Litigious, Strong Modal, Weak Modal, Constraining).
  * Computes Total Words, category word counts, Positive-to-Negative (P/N) Ratio,
    and category percentage shares.
  * Adds a professionally styled 'Sentiment Analysis' tab to the Excel workbook.
- Step 8: Historical Stock Price Retrieval:
  * Identifies the company's actual balance sheet date (reportDate) from SEC filings.
  * Uses yfinance to retrieve historical split-adjusted closing stock prices.
  * Automatically rolls back to the most recent previous trading day for non-trading
    dates (weekends and market holidays) without ever using future dates.
  * Adds a professionally styled 'Stock Prices' tab to the existing Excel workbook.
- Step 9: Final Financial Ratios, Sentiment & Stock Price Analysis:
  * Computes four key ratios: Current Ratio (Liquidity), Debt-to-Assets (Solvency),
    Total Asset Turnover (Efficiency, utilizing Average Total Assets), and Net Profit
    Margin (Profitability).
  * Generates the 'Analysis' tab in the Excel workbook arranged chronologically
    (oldest to newest) with full formula documentation and professional styling.
  * Embeds an openpyxl dual-axis Combination Line Chart comparing Historical Stock
    Price ($) on the primary axis vs. MD&A Sentiment P/N Ratio on the secondary axis.
  * Displays terminal summaries with complete underlying account verification.
"""

import csv
from datetime import datetime, timedelta
import html
import math
import os
from pathlib import Path
import re
import time
import unicodedata
import warnings
import requests
from bs4 import BeautifulSoup, XMLParsedAsHTMLWarning
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference

# Suppress harmless BeautifulSoup XML/HTML parser warnings for SEC inline XBRL
warnings.filterwarnings("ignore", category=XMLParsedAsHTMLWarning)

# ==============================================================================
# PROJECT ROOT DIRECTORY & ENVIRONMENT CONFIGURATION
# ==============================================================================
PROJECT_ROOT = Path(__file__).resolve().parent

# Load local environment variables from .env if python-dotenv is installed
try:
    from dotenv import load_dotenv
    load_dotenv(PROJECT_ROOT / ".env")
except ImportError:
    pass

# ==============================================================================
# SEC EDGAR CONFIGURATION & FAIR ACCESS COMPLIANCE
# ==============================================================================

_SEC_PLACEHOLDER_VALUES = {
    "your name",
    "your university or organization",
    "your_email@example.com",
    "your_email@university.edu",
    "academic research institution",
    "financial research analyst",
    "research_user@example.edu",
    "example.com",
}


def get_sec_credentials():
    """
    Retrieve and validate SEC EDGAR contact information complying with the SEC
    Fair Access policy (https://www.sec.gov/os/accessing-edgar-data).

    The SEC requires all automated requests to declare a User-Agent header
    containing the user's name, organization, and a valid email address.
    Placeholder or fake contact values are strictly rejected.

    Variables are read from:
      1. Environment variables / local '.env' file:
         - SEC_USER_NAME
         - SEC_USER_ORGANIZATION
         - SEC_USER_EMAIL
      2. Streamlit secrets (.streamlit/secrets.toml) if running in Streamlit.

    Returns:
        tuple: (headers_dict, None) if all required variables are set and valid,
               or (None, error_message) if any required variable is missing or placeholder.
    """
    org = os.getenv("SEC_USER_ORGANIZATION")
    name = os.getenv("SEC_USER_NAME")
    email = os.getenv("SEC_USER_EMAIL")

    try:
        import streamlit as st
        if hasattr(st, "secrets"):
            def _get_sec_val(key):
                try:
                    if key in st.secrets:
                        return st.secrets[key]
                    if hasattr(st.secrets, "get") and isinstance(st.secrets.get("sec"), dict):
                        return st.secrets["sec"].get(key)
                except Exception:
                    pass
                return None

            if not org or org.strip().lower() in _SEC_PLACEHOLDER_VALUES:
                val = _get_sec_val("SEC_USER_ORGANIZATION")
                if val:
                    org = val
            if not name or name.strip().lower() in _SEC_PLACEHOLDER_VALUES:
                val = _get_sec_val("SEC_USER_NAME")
                if val:
                    name = val
            if not email or email.strip().lower() in _SEC_PLACEHOLDER_VALUES:
                val = _get_sec_val("SEC_USER_EMAIL")
                if val:
                    email = val
    except Exception:
        pass

    missing = []
    org_val = (org or "").strip()
    name_val = (name or "").strip()
    email_val = (email or "").strip()

    if not org_val or org_val.lower() in _SEC_PLACEHOLDER_VALUES:
        missing.append("SEC_USER_ORGANIZATION")
    if not name_val or name_val.lower() in _SEC_PLACEHOLDER_VALUES:
        missing.append("SEC_USER_NAME")
    if not email_val or email_val.lower() in _SEC_PLACEHOLDER_VALUES or "@" not in email_val:
        missing.append("SEC_USER_EMAIL")

    if missing:
        missing_str = ", ".join(missing)
        error_msg = (
            f"SEC EDGAR User-Agent contact information is not properly configured.\n"
            f"Missing or incomplete required variable(s): {missing_str}\n\n"
            f"Per SEC Fair Access policy, automated requests must declare a verified contact header.\n"
            f"Generic placeholder or fake contact values are not permitted.\n\n"
            f"To configure your contact details:\n"
            f"  • Local CLI (.env): Copy '.env.example' to '.env' and enter your real details:\n"
            f"      SEC_USER_NAME=Your Name\n"
            f"      SEC_USER_ORGANIZATION=Your University or Organization\n"
            f"      SEC_USER_EMAIL=your_email@example.com\n\n"
            f"  • Streamlit Cloud (Secrets): In App Settings → Secrets, configure:\n"
            f"      SEC_USER_ORGANIZATION = \"Your University or Organization\"\n"
            f"      SEC_USER_NAME = \"Your Name\"\n"
            f"      SEC_USER_EMAIL = \"your_email@example.com\"\n\n"
            f"Save the configuration and rerun the application."
        )
        return None, error_msg

    headers = {
        "User-Agent": f"{org_val} - {name_val} ({email_val})",
        "Accept-Encoding": "gzip, deflate"
    }
    return headers, None


def get_sec_headers():
    """
    Return SEC EDGAR request headers if valid credentials are configured,
    or None if any required SEC variable is missing.
    """
    headers, _ = get_sec_credentials()
    return headers


# Politeness rate limiting for SEC EDGAR (max 10 requests/second)
SEC_REQUEST_DELAY = 0.15


# ==============================================================================
# STEP 2 FUNCTIONS: USER INPUT & VALIDATION
# ==============================================================================

def get_ticker_symbol():
    """
    Prompt the user to enter a stock ticker symbol.

    The input is cleaned by stripping whitespace and converting to uppercase.
    Validates that the ticker is non-empty and contains valid characters.

    Returns:
        str: A validated, uppercase ticker symbol.
    """
    while True:
        user_input = input("Enter a public company ticker symbol (e.g., AAPL, MSFT): ")
        ticker = user_input.strip().upper()

        if not ticker:
            print("Error: Ticker symbol cannot be empty. Please try again.\n")
            continue

        # Allow letters as well as '.' or '-' for dual-class shares (e.g., BRK.B)
        cleaned_ticker = ticker.replace(".", "").replace("-", "")
        if not cleaned_ticker.isalpha():
            print("Error: Ticker symbols should only contain letters. Please try again.\n")
            continue

        return ticker


def get_filing_count():
    """
    Prompt the user to enter the number of 10-K filings to retrieve.

    Validates that the input is a whole integer between 1 and 10.

    Returns:
        int: The number of filings requested (1 through 10).
    """
    while True:
        user_input = input("Enter the number of recent 10-K filings to retrieve (1 to 10): ").strip()

        try:
            filing_count = int(user_input)
        except ValueError:
            print("Error: Please enter a valid whole number (integer).\n")
            continue

        if 1 <= filing_count <= 10:
            return filing_count
        else:
            print(f"Error: {filing_count} is out of range. Please choose a number from 1 to 10.\n")


# ==============================================================================
# STEP 3 FUNCTIONS: SEC EDGAR RETRIEVAL & STORAGE
# ==============================================================================

def get_company_cik(ticker):
    """
    Lookup a company's SEC Central Index Key (CIK) using its ticker symbol.

    Args:
        ticker (str): Company stock ticker symbol.

    Returns:
        tuple: (cik_int, company_title) if found, or (None, None) if not found.
    """
    headers, cred_err = get_sec_credentials()
    if not headers:
        print(f"\n[SEC Configuration Error] Cannot connect to SEC EDGAR:\n{cred_err}")
        return None, None

    url = "https://www.sec.gov/files/company_tickers.json"

    try:
        response = requests.get(url, headers=headers, timeout=15)
        response.raise_for_status()
        tickers_data = response.json()
    except requests.exceptions.RequestException as error:
        print(f"\nError connecting to SEC EDGAR company directory: {error}")
        return None, None

    normalized_ticker = ticker.replace(".", "").replace("-", "").upper()

    for item in tickers_data.values():
        entry_ticker = item.get("ticker", "").replace(".", "").replace("-", "").upper()
        if entry_ticker == normalized_ticker:
            cik = item.get("cik_str")
            title = item.get("title")
            return cik, title

    return None, None


def collect_10k_entries(filings_block):
    """
    Pull Form 10-K entries out of one block of SEC submissions data.

    Args:
        filings_block (dict): Column arrays ('form', 'accessionNumber',
                              'filingDate', 'reportDate', 'primaryDocument') from the SEC API.

    Returns:
        list of dict: Metadata for each Form 10-K in the block, newest first.
    """
    forms = filings_block.get("form", [])
    accession_numbers = filings_block.get("accessionNumber", [])
    filing_dates = filings_block.get("filingDate", [])
    report_dates = filings_block.get("reportDate", [])
    primary_documents = filings_block.get("primaryDocument", [])

    entries = []
    for index, form_name in enumerate(forms):
        if form_name == "10-K":
            filing_date = filing_dates[index]
            # Retrieve the company's actual balance sheet date from SEC reportDate
            report_date = (
                report_dates[index]
                if index < len(report_dates) and report_dates[index]
                else filing_date
            )
            entries.append({
                "accession_number": accession_numbers[index],
                "filing_date": filing_date,
                "report_date": report_date,
                "year": filing_date.split("-")[0],
                # Fiscal year = year of the balance sheet date. A December year-end
                # company (e.g., TSLA) files its FY2025 10-K in early 2026.
                "fiscal_year": report_date.split("-")[0],
                "primary_document": primary_documents[index]
            })
    return entries


def get_recent_10k_metadata(cik, count):
    """
    Fetch metadata for the most recent Form 10-K filings for a given CIK.

    Args:
        cik (int): The company's Central Index Key.
        count (int): Number of recent 10-K filings requested.

    Returns:
        list of dict: List containing metadata for each filing.
    """
    cik_padded = str(cik).zfill(10)
    url = f"https://data.sec.gov/submissions/CIK{cik_padded}.json"

    headers, cred_err = get_sec_credentials()
    if not headers:
        print(f"\n[SEC Configuration Error] Cannot retrieve filings:\n{cred_err}")
        return []

    try:
        response = requests.get(url, headers=headers, timeout=15)
        response.raise_for_status()
        submissions = response.json()
    except requests.exceptions.RequestException as error:
        print(f"\nError retrieving filing submissions list from SEC: {error}")
        return []

    available_10ks = collect_10k_entries(submissions.get("filings", {}).get("recent", {}))

    # The 'recent' block only holds the latest ~1,000 filings of any form type.
    # Frequent filers (e.g., MSFT) keep older 10-Ks in additional history pages.
    for page in submissions.get("filings", {}).get("files", []):
        if len(available_10ks) >= count:
            break
        page_url = f"https://data.sec.gov/submissions/{page.get('name')}"
        try:
            time.sleep(SEC_REQUEST_DELAY)
            response = requests.get(page_url, headers=headers, timeout=15)
            response.raise_for_status()
            available_10ks.extend(collect_10k_entries(response.json()))
        except requests.exceptions.RequestException as error:
            print(f"\nWarning: Could not retrieve older filing history page: {error}")
            break

    total_available = len(available_10ks)
    if total_available == 0:
        print(f"\nNotice: No Form 10-K filings were found for this entity.")
        print("(Note for Accounting: Foreign companies such as Sony/SONY file Form 20-F "
              "annual reports instead of domestic Form 10-K).")
        return []

    if total_available < count:
        print(f"\nNotice: You requested {count} filing(s), but only {total_available} "
              f"Form 10-K filing(s) are available.")
        print(f"Downloading all {total_available} available filing(s).")
        return available_10ks

    return available_10ks[:count]


def setup_filing_directory(ticker):
    """
    Create the project folder structure: 'filings/<TICKER>/'.

    Args:
        ticker (str): The company's ticker symbol.

    Returns:
        str: Absolute path to the company's filing folder, or None if failed.
    """
    target_folder = str(PROJECT_ROOT / "filings" / ticker.upper())

    try:
        os.makedirs(target_folder, exist_ok=True)
        return target_folder
    except OSError as error:
        print(f"\nError creating folder '{target_folder}': {error}")
        return None


def download_10k_filings(ticker, cik, filings_metadata, target_folder):
    """
    Download Form 10-K filings into the company folder and rename them.

    Naming convention: <TICKER>_<YEAR>_10-K.html.
    Skips downloading files that already exist on disk.

    Args:
        ticker (str): Company ticker symbol.
        cik (int): Central Index Key.
        filings_metadata (list): List of dictionaries containing filing details.
        target_folder (str): Directory path where files will be saved.

    Returns:
        list of dict: Records with filing date, file name, status, and path.
    """
    download_results = []
    print(f"\nChecking and downloading {len(filings_metadata)} filing(s)...")

    for filing in filings_metadata:
        filing_date = filing["filing_date"]
        year = filing["year"]
        primary_doc = filing["primary_document"]
        accession_no = filing["accession_number"]

        standard_filename = f"{ticker}_{year}_10-K.html"
        save_path = os.path.join(target_folder, standard_filename)

        if os.path.exists(save_path):
            print(f"  [SKIPPED] {standard_filename} already exists.")
            download_results.append({
                "date": filing_date,
                "filename": standard_filename,
                "status": "Already exists (Skipped download)",
                "path": save_path
            })
            continue

        accession_clean = accession_no.replace("-", "")
        doc_url = f"https://www.sec.gov/Archives/edgar/data/{cik}/{accession_clean}/{primary_doc}"

        print(f"  [DOWNLOADING] {standard_filename} from SEC EDGAR...")

        headers, cred_err = get_sec_credentials()
        if not headers:
            print(f"  [FAILED] SEC contact credentials not configured:\n{cred_err}")
            download_results.append({
                "date": filing_date,
                "filename": standard_filename,
                "status": "Download failed: SEC credentials not configured",
                "path": save_path
            })
            continue

        try:
            time.sleep(SEC_REQUEST_DELAY)
            response = requests.get(doc_url, headers=headers, timeout=30)
            response.raise_for_status()

            with open(save_path, "wb") as file_out:
                file_out.write(response.content)

            download_results.append({
                "date": filing_date,
                "filename": standard_filename,
                "status": "Downloaded successfully",
                "path": save_path
            })

        except requests.exceptions.RequestException as net_error:
            print(f"  [FAILED] Could not download {standard_filename}: {net_error}")
            download_results.append({
                "date": filing_date,
                "filename": standard_filename,
                "status": f"Download failed: {net_error}",
                "path": save_path
            })
        except OSError as file_error:
            print(f"  [FAILED] Could not write file {standard_filename}: {file_error}")
            download_results.append({
                "date": filing_date,
                "filename": standard_filename,
                "status": f"Save failed: {file_error}",
                "path": save_path
            })

    return download_results


def display_download_confirmation(company_name, ticker, cik, target_folder, results):
    """
    Display a summary of downloaded and existing filings.
    """
    print("\n" + "=" * 70)
    print("                     SEC 10-K RETRIEVAL SUMMARY")
    print("=" * 70)
    print(f"Company Name   : {company_name}")
    print(f"Ticker Symbol  : {ticker}")
    print(f"SEC CIK Number : {str(cik).zfill(10)}")
    print(f"Saved Directory: {target_folder}")
    print("-" * 70)
    print(f"{'Filing Date':<14} {'File Name':<28} {'Status'}")
    print("-" * 70)

    for item in results:
        print(f"{item['date']:<14} {item['filename']:<28} {item['status']}")

    print("=" * 70)


# ==============================================================================
# STEP 4 FUNCTIONS: MD&A EXTRACTION & TEXT PROCESSING
# ==============================================================================

def setup_extracted_text_directory(ticker):
    """
    Create the storage folder for extracted plain text: 'extracted_text/<TICKER>/'.

    Args:
        ticker (str): Company ticker symbol.

    Returns:
        str: Absolute path to the extraction directory, or None if failed.
    """
    target_folder = str(PROJECT_ROOT / "extracted_text" / ticker.upper())

    try:
        os.makedirs(target_folder, exist_ok=True)
        return target_folder
    except OSError as error:
        print(f"\nError creating extraction directory '{target_folder}': {error}")
        return None


def clean_html_to_text(html_content):
    """
    Parse raw 10-K HTML content into clean text while preserving paragraph breaks.

    Removes script/style blocks and unwraps inline styling tags (such as font, span, b)
    so that words split across tags inside section headings remain contiguous.

    Args:
        html_content (str): The raw HTML string of the 10-K filing.

    Returns:
        str: Cleaned text with normalized whitespace and paragraph spacing.
    """
    cleaned_html = html_content.replace("&#160;", " ").replace("&nbsp;", " ")
    cleaned_html = cleaned_html.replace("&#8217;", "'").replace("’", "'")
    cleaned_html = cleaned_html.replace("&#8220;", '"').replace("&#8221;", '"')

    soup = BeautifulSoup(cleaned_html, "html.parser")

    for element in soup(["script", "style", "header", "footer"]):
        element.decompose()

    for inline_tag in soup.find_all(["font", "span", "b", "strong", "i", "em", "u", "a"]):
        inline_tag.unwrap()

    raw_text = soup.get_text(separator="\n")
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    return "\n".join(lines)


def extract_mda_section(clean_text):
    """
    Locate and extract the MD&A section from cleaned 10-K text.

    Finds Item 7 ("Management's Discussion and Analysis..."), filters out
    Table of Contents references, and stops before Item 7A or Item 8.

    Args:
        clean_text (str): Cleaned text produced by clean_html_to_text.

    Returns:
        tuple: (mda_text, word_count) if found, or (None, 0) if not found.
    """
    start_pattern = re.compile(
        r"(?:I[\s\n]*TEM|Item)\s*7\s*[\.\:\-—]?\s*(?:\n\s*)?(?:MANAGEMENT[\'’]S|Management[\'’]s)\s+Discussion\s+and\s+Analysis",
        re.IGNORECASE
    )

    matches_start = list(start_pattern.finditer(clean_text))

    if not matches_start:
        fallback_start = re.compile(
            r"MANAGEMENT[\'’]S\s+DISCUSSION\s+AND\s+ANALYSIS\s+OF\s+FINANCIAL\s+CONDITION",
            re.IGNORECASE
        )
        matches_start = list(fallback_start.finditer(clean_text))

    if not matches_start:
        return None, 0

    valid_start = None
    for match in matches_start:
        pos = match.start()
        lookahead = clean_text[pos:pos + 600]
        if re.search(r"\b(?:I[\s\n]*TEM|Item)\s*(?:7A|8)\b", lookahead, re.IGNORECASE):
            continue
        valid_start = pos
        break

    if valid_start is None:
        valid_start = matches_start[-1].start()

    end_pattern = re.compile(
        r"(?:I[\s\n]*TEM|Item)\s*(?:7A|8)\s*[\.\:\-—]?\s*(?:\n\s*)?(?:QUANTITATIVE|Quantitative|FINANCIAL|Financial)",
        re.IGNORECASE
    )

    end_matches = [m for m in end_pattern.finditer(clean_text) if m.start() > valid_start]

    if not end_matches:
        fallback_end = re.compile(r"^\s*(?:I[\s\n]*TEM|Item)\s*(?:7A|8)\b", re.MULTILINE | re.IGNORECASE)
        end_matches = [m for m in fallback_end.finditer(clean_text) if m.start() > valid_start]

    if end_matches:
        valid_end = end_matches[0].start()
        mda_raw = clean_text[valid_start:valid_end].strip()
    else:
        valid_end = min(len(clean_text), valid_start + 150000)
        mda_raw = clean_text[valid_start:valid_end].strip()

    paragraphs = [p.strip() for p in mda_raw.splitlines() if p.strip()]
    formatted_mda = "\n\n".join(paragraphs)

    word_count = len(formatted_mda.split())
    return formatted_mda, word_count


def process_mda_extractions(ticker, filings_folder, extracted_folder, filings_metadata=None):
    """
    Process downloaded 10-K HTML filings for a company and extract MD&A text.

    Only processes the filings associated with the current user request (from filings_metadata).
    Saves each extracted section as '<TICKER>_<YEAR>_MDA.txt'.
    Skips extraction if the output file already exists.

    Args:
        ticker (str): Company ticker symbol.
        filings_folder (str): Directory containing source 10-K HTML files.
        extracted_folder (str): Directory where extracted text files will be stored.
        filings_metadata (list, optional): Metadata list for the current request.

    Returns:
        list of dict: Detailed extraction report records.
    """
    extraction_results = []

    if not os.path.exists(filings_folder):
        print(f"\nError: Filings folder '{filings_folder}' not found.")
        return extraction_results

    # Filter strictly to the filings requested in this session if provided
    if filings_metadata:
        target_names = [f"{ticker}_{item['year']}_10-K.html" for item in filings_metadata]
        filing_files = [f for f in target_names if os.path.exists(os.path.join(filings_folder, f))]
    else:
        filing_files = sorted([f for f in os.listdir(filings_folder) if f.endswith(".html")])

    if not filing_files:
        print(f"\nNotice: No downloaded HTML filings found in '{filings_folder}'.")
        return extraction_results

    print(f"\nProcessing MD&A extraction for {len(filing_files)} filing(s)...")

    for html_file in filing_files:
        source_path = os.path.join(filings_folder, html_file)
        file_parts = html_file.replace(".html", "").split("_")
        year = file_parts[1] if len(file_parts) >= 2 else "UNKNOWN"

        output_filename = f"{ticker}_{year}_MDA.txt"
        output_path = os.path.join(extracted_folder, output_filename)

        if os.path.exists(output_path):
            try:
                with open(output_path, "r", encoding="utf-8") as existing_file:
                    content = existing_file.read()
                    existing_words = len(content.split())
            except OSError:
                existing_words = 0

            print(f"  [SKIPPED] {output_filename} already exists.")
            extraction_results.append({
                "source_file": html_file,
                "found": True,
                "output_file": output_filename,
                "save_path": output_path,
                "words": existing_words,
                "status": "Already extracted (Skipped)"
            })
            continue

        print(f"  [EXTRACTING] MD&A from {html_file}...")

        try:
            with open(source_path, "r", encoding="utf-8", errors="ignore") as file_in:
                html_data = file_in.read()
        except OSError as read_error:
            print(f"  [ERROR] Inability to read filing {html_file}: {read_error}")
            extraction_results.append({
                "source_file": html_file,
                "found": False,
                "output_file": output_filename,
                "save_path": output_path,
                "words": 0,
                "status": f"Read error: {read_error}"
            })
            continue

        clean_text = clean_html_to_text(html_data)
        mda_text, word_count = extract_mda_section(clean_text)

        if not mda_text or word_count == 0:
            print(f"  [WARNING] MD&A section not found in {html_file}.")
            extraction_results.append({
                "source_file": html_file,
                "found": False,
                "output_file": output_filename,
                "save_path": output_path,
                "words": 0,
                "status": "MD&A section not found"
            })
            continue

        try:
            with open(output_path, "w", encoding="utf-8") as file_out:
                file_out.write(mda_text)

            print(f"  [SAVED] {output_filename} ({word_count:,} words)")
            extraction_results.append({
                "source_file": html_file,
                "found": True,
                "output_file": output_filename,
                "save_path": output_path,
                "words": word_count,
                "status": "Extracted & Saved"
            })

        except OSError as write_error:
            print(f"  [ERROR] Could not save extracted text for {html_file}: {write_error}")
            extraction_results.append({
                "source_file": html_file,
                "found": True,
                "output_file": output_filename,
                "save_path": output_path,
                "words": word_count,
                "status": f"Save error: {write_error}"
            })

    return extraction_results


def display_mda_summary(ticker, extracted_folder, results):
    """
    Display a summary of the extracted MD&A sections.
    """
    print("\n" + "=" * 80)
    print("                     STEP 4: MD&A EXTRACTION SUMMARY")
    print("=" * 80)
    print(f"Company Ticker Symbol : {ticker}")
    print(f"Saved Directory Path  : {extracted_folder}")
    print("-" * 80)
    print(f"{'Source 10-K File':<24} {'MD&A Found':<12} {'Output Text File':<22} {'Words':<9} {'Status'}")
    print("-" * 80)

    for record in results:
        found_str = "Yes" if record["found"] else "No"
        words_str = f"{record['words']:,}" if record["words"] > 0 else "-"
        print(f"{record['source_file']:<24} {found_str:<12} {record['output_file']:<22} {words_str:<9} {record['status']}")

    print("=" * 80)


# ==============================================================================
# STEP 5 FUNCTIONS: FINANCIAL STATEMENTS EXTRACTION (EXCEL WORKBOOK)
# ==============================================================================

def setup_output_directory():
    """
    Create the 'output' directory inside the project folder if it doesn't exist.

    Returns:
        str: Absolute path to the output directory, or None if failed.
    """
    output_dir = str(PROJECT_ROOT / "output")
    try:
        os.makedirs(output_dir, exist_ok=True)
        return output_dir
    except OSError as err:
        print(f"\nError creating output directory '{output_dir}': {err}")
        return None


def get_item8_slice(html_content):
    """
    Isolate the Item 8 section of the 10-K filing to avoid picking up
    summary tables from Item 7 (MD&A) or footnote details elsewhere.

    Args:
        html_content (str): Full HTML content of the 10-K.

    Returns:
        str: The HTML slice containing Item 8 financial statements.
    """
    html_norm = html_content.replace("&#160;", " ").replace("&nbsp;", " ")

    # Older filings split headings across tags, e.g. 'Item 8.</span></td><td>Financial
    # Statements', so allow whitespace and HTML tags between the heading words.
    gap = r"(?:\s|<[^>]+>)*"

    # Find beginning of Item 8: Financial Statements
    # Stop at 'Financial': some filings even split 'STATE</span><span>MENTS' with tags
    # 'I{gap}TEM' also covers small-caps headings written as 'I</font><font>TEM 8.'
    p8 = re.compile(
        rf"I{gap}TEM{gap}8{gap}[\.\:\-—]?{gap}Financial",
        re.IGNORECASE
    )
    matches_8 = list(p8.finditer(html_norm))
    if not matches_8:
        return html_norm

    # Find beginning of Item 9 / 9A: Changes in ... / Controls and Procedures
    p9 = re.compile(
        rf"I{gap}TEM{gap}9A?{gap}[\.\:\-—]?{gap}(?:Changes|Controls)",
        re.IGNORECASE
    )
    positions_9 = [m.start() for m in p9.finditer(html_norm)]

    # The Table of Contents also lists Item 8 followed closely by Item 9, so pick
    # the Item 8 heading followed by the most content before the next Item 9.
    best_slice = ""
    for match in matches_8:
        pos8 = match.start()
        later_9 = [p for p in positions_9 if p > pos8]
        # Fallback: slice up to 2 million characters after Item 8 start
        end = later_9[0] if later_9 else pos8 + 2000000
        if end - pos8 > len(best_slice):
            best_slice = html_norm[pos8:end]

    # Only a Table of Contents entry was found; search the whole filing instead
    if len(best_slice) < 50000:
        return html_norm

    return best_slice


def find_statement_tables(item8_html):
    """
    Scan HTML tables within Item 8 and identify the primary financial statements:
    - Income Statement (Operations, Income, or Earnings)
    - Balance Sheet (Financial Position or Balance Sheets)
    - Statement of Cash Flows

    Args:
        item8_html (str): The HTML segment containing Item 8.

    Returns:
        dict: Mapping statement keys ('income', 'balance', 'cash') to BeautifulSoup table elements.
    """
    soup = BeautifulSoup(item8_html, "lxml")
    tables = soup.find_all("table")
    statement_tables = {}

    for table in tables:
        text = table.get_text(separator=" ", strip=True)
        if len(text) < 150:
            continue

        # Skip small summary tables (e.g., a 3-row cash flow summary in MD&A)
        if len(table.find_all("tr")) < 8:
            continue

        # Skip Table of Contents / Index tables
        if "Index to Consolidated Financial Statements" in text or "Report of Independent" in text:
            continue

        text_lower = text.lower()

        # 1. Income Statement Detection
        if "income" not in statement_tables:
            has_inc_rev = any(k in text_lower for k in [
                "revenue", "revenues", "net sales", "total net sales", "total revenue"
            ])
            has_inc_op = any(k in text_lower for k in [
                "operating income", "income from operations", "gross margin", "cost of sales", "cost of revenue"
            ])
            has_inc_net = any(k in text_lower for k in [
                "net income", "net earnings", "net loss"
            ])
            if has_inc_rev and has_inc_op and has_inc_net:
                if "comprehensive" not in text_lower[:120] and "other comprehensive" not in text_lower[:120]:
                    statement_tables["income"] = table

        # 2. Balance Sheet Detection
        if "balance" not in statement_tables:
            has_assets = any(k in text_lower for k in ["total assets", "current assets"])
            has_liab = any(k in text_lower for k in [
                "total liabilities", "current liabilities", "stockholders’ equity",
                "stockholders' equity", "shareholders’ equity", "shareholders' equity"
            ])
            if has_assets and has_liab:
                statement_tables["balance"] = table

        # 3. Statement of Cash Flows Detection
        if "cash" not in statement_tables:
            has_cf_op = any(k in text_lower for k in [
                "operating activities", "cash from operations", "cash flows from operating",
                "operations net income", "net cash from operations"
            ])
            has_cf_inv_fin = any(k in text_lower for k in [
                "investing activities", "cash used in investing", "cash flows from investing",
                "net cash used in financing", "financing activities"
            ])
            if has_cf_op and has_cf_inv_fin:
                statement_tables["cash"] = table

    return statement_tables


def parse_statement_table(table):
    """
    Extract structured accounting rows and values from an HTML financial statement table.

    Handles SEC-specific formatting challenges:
    - Extracts fiscal years from date headers (e.g., 'September 27, 2025' -> '2025').
    - Discovers statement unit scale (e.g., 'In millions, except per share amounts').
    - Converts parentheses '(12,345)' into negative floating-point numbers.
    - Strips '$' currency markers and ',' separators.
    - Converts em dashes '—' or '-' into 0.0.
    - Preserves row section headers (e.g. 'Operating expenses:').

    Args:
        table (bs4.element.Tag): The HTML table element.

    Returns:
        tuple: (years_list, unit_scale_str, parsed_line_items)
               where parsed_line_items is a list of tuples: (label, values_dict, is_header).
    """
    rows = table.find_all("tr")
    years = []
    unit_str = ""
    header_row_idx = -1

    # Search preceding sibling tags for scale / unit declaration
    curr = table
    for _ in range(15):
        curr = curr.previous_element
        if not curr:
            break
        if getattr(curr, "name", None) in ["div", "p", "span"]:
            txt = curr.get_text(strip=True)
            if any(w in txt.lower() for w in ["in millions", "in thousands", "dollars in millions"]):
                unit_str = txt
                break

    # Identify fiscal-year columns by inspecting the first 8 rows
    for idx, r in enumerate(rows[:8]):
        text = r.get_text(separator=" ", strip=True)

        if any(w in text.lower() for w in ["in millions", "in thousands", "dollars in millions"]):
            unit_str = text

        found_years = re.findall(r"\b(20\d\d)\b", text)
        unique_years = []
        for y in found_years:
            if y not in unique_years:
                unique_years.append(y)

        # Most annual statements report 2 or 3 comparative fiscal years
        if len(unique_years) >= 2:
            years = unique_years
            header_row_idx = idx
            break

    if not years:
        return None, unit_str, []

    line_items = []
    for r in rows[header_row_idx + 1:]:
        cells = [c.get_text(strip=True) for c in r.find_all(["td", "th"])]
        non_empty = [c for c in cells if c]
        if not non_empty:
            continue

        label = ""
        val_tokens = []
        for cell_txt in non_empty:
            if not label:
                # The first cell containing word characters (and not currency codes) is the line item
                if re.search(r"[a-zA-Z]{2,}", cell_txt) and cell_txt not in ["N/A", "USD"]:
                    label = cell_txt
                    continue
            val_tokens.append(cell_txt)

        if not label:
            continue

        # Clean label: remove footnote markers like (1) or trailing colon
        clean_label = re.sub(r"\s*\(\d+\)$", "", label).strip()

        # Join tokens and extract numbers with negative parentheses or em dashes
        token_str = " ".join(val_tokens).replace("$", " ")
        num_matches = re.findall(
            r"(?:\(\s*[\d,]+(?:\.\d+)?\s*\)|-?\s*[\d,]+(?:\.\d+)?|[—–\-])",
            token_str
        )

        parsed_vals = []
        for match in num_matches:
            m_clean = match.strip()
            if m_clean in ["—", "–", "-"]:
                parsed_vals.append(0.0)
            elif m_clean.startswith("(") and m_clean.endswith(")"):
                num_part = m_clean[1:-1].replace(",", "").strip()
                try:
                    parsed_vals.append(-float(num_part))
                except ValueError:
                    pass
            else:
                num_part = m_clean.replace(",", "").replace(" ", "")
                try:
                    parsed_vals.append(float(num_part))
                except ValueError:
                    pass

        # Valid numerical data row matching column count
        if len(parsed_vals) == len(years):
            row_dict = {years[i]: parsed_vals[i] for i in range(len(years))}
            line_items.append((clean_label, row_dict, False))
        elif len(parsed_vals) == 0:
            # Section header row (e.g., "Current assets:")
            line_items.append((clean_label, {}, True))

    return years, unit_str, line_items


def unit_multiplier(unit_text):
    """
    Read the reporting scale from a statement's unit note.

    Uses the first scale word found, so '(In millions, except number of shares,
    which are reflected in thousands)' is treated as millions.

    Args:
        unit_text (str): Unit declaration such as '(in thousands, except per share data)'.

    Returns:
        float or None: 1e3, 1e6, or 1e9, or None if no scale is stated.
    """
    match = re.search(r"thousands|millions|billions", (unit_text or "").lower())
    if not match:
        return None
    return {"thousands": 1e3, "millions": 1e6, "billions": 1e9}[match.group(0)]


def normalize_label(label):
    """
    Normalize a line-item label for matching across sections and filings.

    Lowercases, unifies apostrophes, and strips digits and punctuation so that
    labels containing changing share counts (e.g., '15,116,786 shares issued')
    still match the same row in other years.

    Args:
        label (str): Raw line-item label.

    Returns:
        str: Normalized matching key.
    """
    text = label.lower().replace("’", "'")
    text = re.sub(r"[\d$.,;:()/'\-–—]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


# Line items that companies rename between years, mapped to one canonical name.
# Each rule is (pattern on the normalized label, canonical name, section hint).
# The section hint ('current' / 'non current') separates items like current vs.
# non-current term debt when an older filing labels them differently.
LABEL_SYNONYMS = [
    (r"^short term marketable securities", "marketable securities", "current"),
    (r"^long term marketable securities", "marketable securities", "non current"),
    (r"^current portion of (?:long term |term )?debt", "term debt", "current"),
    (r"^(?:long term debt|non current portion of term debt)$", "term debt", "non current"),
    (r"^deferred revenue non current$", "deferred revenue", "non current"),
    (r"^accounts receivable", "accounts receivable", None),
    (r"^(?:total )?net sales$", "total net sales", None),
    (r"^(?:total )?cost of sales$", "total cost of sales", None),
    (r"^(?:retained earnings|accumulated deficit)", "retained earnings", None),
    (r"^accumulated other comprehensive", "accumulated other comprehensive income", None),
    (r"^deferred income tax", "deferred income tax", None),
    (r"commercial paper net$", "commercial paper net", None),
    (r"^(?:net )?cash\b.*\boperating activities$", "net cash operating activities", None),
    (r"^(?:net )?cash\b.*\binvesting activities$", "net cash investing activities", None),
    (r"^(?:net )?cash\b.*\bfinancing activities$", "net cash financing activities", None),
    (r"^cash\b.*\bbeginning", "cash beginning balance", None),
    (r"^cash\b.*\b(?:ending balances?|end of)", "cash ending balance", None),
    (r"^increase(?: decrease)? in cash", "increase decrease in cash", None),
]


def canonical_label(label):
    """
    Map a line-item label to a canonical name used to merge renamed rows.

    Args:
        label (str): Raw line-item label.

    Returns:
        tuple: (canonical_name, section_hint) where section_hint is
               'current', 'non current', or None.
    """
    # Ignore parentheticals such as '(gains)' or '(loss)' when matching
    text = normalize_label(re.sub(r"\([^)]*\)", " ", label))
    for pattern, canonical, hint in LABEL_SYNONYMS:
        if re.search(pattern, text):
            return canonical, hint
    return text, None


def merge_equivalent_rows(statement, years):
    """
    Fold rows from older filings into matching rows from newer filings.

    Companies rename line items over ten years (e.g., 'Net sales' became
    'Total net sales', 'Long-term debt' became non-current 'Term debt'). A row
    first seen in an older filing is merged into an earlier row when:
      * both labels map to the same canonical name (see LABEL_SYNONYMS), or
      * both rows sit in the same section and report identical values,
    Rows that disagree on a shared fiscal year (a restatement) merge only when
    they are the same item in the same section, keeping the newer filing's value.
    Section headers from older filings left with no data rows are dropped.

    Args:
        statement (dict): Merged statement data with 'items_order',
                          'items_values', and 'items_origin'.
        years (list): Fiscal years that will appear in the workbook.

    Returns:
        list: The merged items_order list.
    """
    values = statement["items_values"]
    origin = statement["items_origin"]
    items = statement["items_order"]
    removed = set()

    def shown(key):
        return {yr: val for yr, val in values.get(key, {}).items() if yr in years}

    # Process newer filings first so each row merges into an already-final row
    for key, label, is_header in sorted(items, key=lambda item: origin.get(item[0], 0)):
        row_vals = shown(key)
        if is_header or origin.get(key, 0) == 0 or not row_vals:
            continue

        canon, hint = canonical_label(label)
        section = key[0]
        best_target, best_rank = None, None

        for other_key, other_label, other_header in items:
            if (other_header or other_key in removed
                    or origin.get(other_key, 0) >= origin.get(key, 0)):
                continue
            other_vals = shown(other_key)
            overlap = [yr for yr in row_vals if yr in other_vals]
            conflict = any(other_vals[yr] != row_vals[yr] for yr in overlap)

            other_canon, _ = canonical_label(other_label)
            other_section = other_key[0]
            same_values_same_section = bool(overlap) and not conflict and other_section == section
            if other_canon != canon and not same_values_same_section:
                continue

            # A current / non-current hint always decides the target section
            if hint and not other_section.startswith(hint):
                continue

            # Rows that disagree on a shared year merge only when they are the same
            # item in the same section (a restatement); the newer filing's value wins
            if conflict and other_section != section:
                continue

            # Prefer the same section, then the section named by the hint
            if other_section == section:
                rank = 0
            elif hint:
                rank = 1
            else:
                rank = 2
            if best_rank is None or rank < best_rank:
                best_target, best_rank = other_key, rank

        if best_target is None:
            continue

        for yr, val in values[key].items():
            values[best_target].setdefault(yr, val)
        removed.add(key)

    kept = [item for item in items if item[0] not in removed]

    # Drop section headers from older filings whose rows were all merged away
    final = []
    for idx, (key, label, is_header) in enumerate(kept):
        if is_header and origin.get(key, 0) > 0:
            following = []
            for next_key, _, next_header in kept[idx + 1:]:
                if next_header:
                    break
                following.append(next_key)
            if not any(shown(k) for k in following):
                continue
        final.append((key, label, is_header))

    return final


def extract_and_merge_statements(filings_folder, ticker=None, filings_metadata=None, max_periods=None):
    """
    Process all downloaded 10-K filings for a company, extract their financial
    statements, and merge multi-year data so that each fiscal year appears only once.

    Args:
        filings_folder (str): Directory containing downloaded 10-K HTML files.

    Returns:
        dict: Aggregated financial statements with structure:
              {
                  'Income Statement': {'years': [...], 'unit': '...', 'items': [...]},
                  'Balance Sheet': {'years': [...], 'unit': '...', 'items': [...]},
                  'Cash Flows': {'years': [...], 'unit': '...', 'items': [...]}
              }
    """
    # Filter strictly to the filings requested in this session if provided
    if filings_metadata and ticker:
        target_names = [f"{ticker}_{item['year']}_10-K.html" for item in filings_metadata]
        html_files = [f for f in target_names if os.path.exists(os.path.join(filings_folder, f))]
        html_files.sort(reverse=True)
    else:
        html_files = sorted([f for f in os.listdir(filings_folder) if f.endswith(".html")], reverse=True)
    if not html_files:
        return {}

    statement_configs = {
        "income": {"sheet_name": "Income Statement", "default_unit": "In millions, except per share amounts"},
        "balance": {"sheet_name": "Balance Sheet", "default_unit": "In millions"},
        "cash": {"sheet_name": "Cash Flows", "default_unit": "In millions"}
    }

    merged_data = {
        cfg["sheet_name"]: {
            "years": set(),
            "unit": "",
            "base_multiplier": None,
            "items_order": [],
            "items_values": {},
            "items_origin": {}
        }
        for cfg in statement_configs.values()
    }

    print(f"\nExtracting financial statements across {len(html_files)} filing(s)...")

    for file_idx, html_file in enumerate(html_files):
        fpath = os.path.join(filings_folder, html_file)
        try:
            with open(fpath, "r", encoding="utf-8", errors="ignore") as f:
                raw_html = f.read()
        except OSError as err:
            print(f"  [WARNING] Could not read filing {html_file}: {err}")
            continue

        # Isolate Item 8 and find tables
        item8_segment = get_item8_slice(raw_html)
        tables_dict = find_statement_tables(item8_segment)

        for stmt_key, cfg in statement_configs.items():
            sheet_title = cfg["sheet_name"]
            if stmt_key not in tables_dict:
                continue

            tbl = tables_dict[stmt_key]
            years, unit_scale, line_items = parse_statement_table(tbl)

            if not years or not line_items:
                continue

            target = merged_data[sheet_title]

            # Adopt the unit from the most recent filing if not already recorded
            if not target["unit"] and unit_scale:
                target["unit"] = unit_scale

            # Older filings may report in thousands while newer ones use millions.
            # Rescale every filing to the most recent filing's unit.
            table_multiplier = unit_multiplier(unit_scale)
            if target["base_multiplier"] is None and table_multiplier:
                target["base_multiplier"] = table_multiplier
            scale_factor = 1.0
            if table_multiplier and target["base_multiplier"]:
                scale_factor = table_multiplier / target["base_multiplier"]

            # Merge years
            for y in years:
                target["years"].add(y)

            # Merge line items. Rows are keyed by (section header, label) so that
            # repeated labels such as 'Products', 'Basic', 'Term debt' or 'Other'
            # in different sections are kept as separate rows.
            section = ""
            seen_keys = {}
            prev_key = None
            for label, vals, is_header in line_items:
                if is_header:
                    section = normalize_label(label)
                    base_key = ("", section, True)
                else:
                    base_key = (section, normalize_label(label), False)

                # Disambiguate a label repeated within the same section of one table
                occurrence = seen_keys.get(base_key, 0)
                seen_keys[base_key] = occurrence + 1
                key = base_key + (occurrence,)

                if key not in target["items_values"]:
                    # Place rows that only older filings report right after the
                    # row that preceded them in that filing, not at the bottom
                    order = target["items_order"]
                    if prev_key is not None:
                        position = next(i for i, item in enumerate(order) if item[0] == prev_key) + 1
                    else:
                        position = 0 if file_idx > 0 else len(order)
                    order.insert(position, (key, label, is_header))
                    target["items_values"][key] = {}
                    target["items_origin"][key] = file_idx
                prev_key = key

                # Per-share amounts are never scaled by the thousands/millions unit
                is_per_share = "per share" in normalize_label(label) or "per share" in section
                factor = 1.0 if is_per_share else scale_factor

                # Fill values for years not already present from newer filings.
                # Rescaled values are rounded to whole units, as the newer
                # statements present them, so matching rows compare equal.
                for yr, val in vals.items():
                    if yr not in target["items_values"][key]:
                        target["items_values"][key][yr] = (
                            float(round(val * factor)) if factor != 1.0 else val
                        )

    # Finalize sorted years and structured results
    final_statements = {}
    for sheet_name, data in merged_data.items():
        sorted_years = sorted(list(data["years"]), reverse=True)
        # Limit to the exact requested number of fiscal periods (max_periods)
        if max_periods and len(sorted_years) > max_periods:
            sorted_years = sorted_years[:max_periods]
        final_unit = data["unit"] if data["unit"] else "In millions, unless otherwise noted"
        final_statements[sheet_name] = {
            "years": sorted_years,
            "unit": final_unit,
            "items": merge_equivalent_rows(data, sorted_years),
            "values": data["items_values"]
        }

    return final_statements


def create_financial_analysis_workbook(company_name, ticker, statement_data, output_filepath):
    """
    Build a professionally styled Excel workbook containing the three financial
    statement tabs side by side across fiscal years.

    Formatting Details:
    - Row 1: Statement Title (Bold, 14pt, Dark Navy Blue: #1F497D)
    - Row 2: Scale/Unit subtitle (Italic, 10pt, Muted Gray: #595959)
    - Row 4: Column Headers (Bold white text on dark blue background)
    - Freeze Panes enabled at Row 5 to keep column headers pinned while scrolling.
    - True Excel accounting number formatting: #,##0;(#,##0);"-"
    - Auto-fitted column widths for presentation quality.

    Args:
        company_name (str): Full company name.
        ticker (str): Stock ticker symbol.
        statement_data (dict): Output from extract_and_merge_statements.
        output_filepath (str): Target .xlsx file destination.

    Returns:
        bool: True if created successfully, False otherwise.
    """
    workbook = openpyxl.Workbook()
    default_sheet = workbook.active

    # Professional Style Palettes
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="595959")
    section_font = Font(name="Calibri", size=10, bold=True, color="000000")
    regular_font = Font(name="Calibri", size=10, color="000000")

    sheet_keys = ["Income Statement", "Balance Sheet", "Cash Flows"]

    for idx, sheet_name in enumerate(sheet_keys):
        # Use existing default sheet for the first tab, create new ones for others
        if idx == 0:
            worksheet = default_sheet
            worksheet.title = sheet_name
        else:
            worksheet = workbook.create_sheet(title=sheet_name)

        data = statement_data.get(sheet_name, {})
        years = data.get("years", [])
        items_order = data.get("items", [])
        values_dict = data.get("values", {})
        unit_str = data.get("unit", "In millions, unless otherwise noted")

        # Row 1: Title Header
        title_text = f"{company_name.upper()} - {sheet_name.upper()}"
        worksheet["A1"] = title_text
        worksheet["A1"].font = title_font

        # Row 2: Unit Subtitle
        worksheet["A2"] = f"Scale / Reporting Currency: {unit_str}"
        worksheet["A2"].font = subtitle_font

        # Row 4: Table Headers
        headers = ["Line Item / Account"] + years
        for col_idx, header_text in enumerate(headers, 1):
            cell = worksheet.cell(row=4, column=col_idx, value=header_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="left" if col_idx == 1 else "right",
                vertical="center"
            )

        # Rows 5+: Data Rows
        current_row = 5
        max_label_len = len("Line Item / Account")

        for key, label, is_header in items_order:
            if len(label) > max_label_len:
                max_label_len = len(label)

            label_cell = worksheet.cell(row=current_row, column=1, value=label)

            if is_header:
                label_cell.font = section_font
            else:
                row_vals = values_dict.get(key, {})
                # Skip line item if it has no values in any of the selected years
                if not any(row_vals.get(yr) is not None for yr in years):
                    continue

                label_cell.font = regular_font
                # Populate numerical columns for each fiscal year
                for yr_idx, yr in enumerate(years, 2):
                    val = row_vals.get(yr)
                    if val is not None:
                        val_cell = worksheet.cell(row=current_row, column=yr_idx, value=val)
                        val_cell.alignment = Alignment(horizontal="right")

                        # Format per-share decimal rows vs whole dollar rows
                        if any(term in label.lower() for term in ["per share", "per-share", "share:"]):
                            val_cell.number_format = "#,##0.00;(#,##0.00);\"-\""
                        else:
                            val_cell.number_format = "#,##0;(#,##0);\"-\""

            current_row += 1

        # Freeze headers at row 5 so title & columns stay visible while scrolling
        worksheet.freeze_panes = "B5"
        worksheet.views.sheetView[0].showGridLines = True

        # Set readable column widths
        worksheet.column_dimensions["A"].width = max(38, min(max_label_len + 5, 68))
        for col_idx in range(2, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            worksheet.column_dimensions[col_letter].width = 16

    try:
        workbook.save(output_filepath)
        return True
    except OSError as save_err:
        print(f"\nError saving Excel file '{output_filepath}': {save_err}")
        return False


def display_step5_summary(ticker, workbook_path, statement_data):
    """
    Display a formatted confirmation table for Step 5 in the terminal.
    """
    print("\n" + "=" * 80)
    print("            STEP 5: FINANCIAL STATEMENTS EXTRACTION SUMMARY")
    print("=" * 80)
    print(f"Company Ticker Symbol : {ticker}")
    print(f"Excel Workbook Path   : {workbook_path}")
    print("-" * 80)
    print(f"{'Worksheet Tab':<22} {'Fiscal Years Included':<28} {'Line Items':<12} {'Status'}")
    print("-" * 80)

    for sheet_name in ["Income Statement", "Balance Sheet", "Cash Flows"]:
        data = statement_data.get(sheet_name, {})
        years = data.get("years", [])
        items_count = len(data.get("items", []))
        years_str = ", ".join(years) if years else "None"
        status_str = "Extracted & Formatted" if items_count > 0 else "Not Found"

        print(f"{sheet_name:<22} {years_str:<28} {items_count:<12} {status_str}")

    print("=" * 80)
    print("Step 5 complete! Structured financial analysis workbook is ready.\n")


# ==============================================================================
# STEP 6: MD&A TEXT CLEANING FOR SENTIMENT ANALYSIS
# ==============================================================================

def setup_cleaned_text_directory(ticker):
    """
    Create a dedicated directory to store cleaned MD&A text files.
    Target directory: cleaned_text/<TICKER>/

    Args:
        ticker (str): Company ticker symbol.

    Returns:
        str: Absolute path to the ticker's cleaned text folder, or None on error.
    """
    ticker_clean_dir = str(PROJECT_ROOT / "cleaned_text" / ticker.upper())

    try:
        os.makedirs(ticker_clean_dir, exist_ok=True)
        return ticker_clean_dir
    except OSError as err:
        print(f"\nError: Could not create cleaned text directory '{ticker_clean_dir}': {err}")
        return None


def clean_mda_text(raw_text):
    """
    Cleans raw MD&A text extracted from 10-K filings to prepare it for
    Loughran-McDonald financial sentiment analysis (Step 7).

    Cleaning rules applied:
    1. Decodes HTML entities and strips any residual HTML tags.
    2. Normalizes Unicode characters via NFKD decomposition.
    3. Converts curly quotation marks and curly apostrophes to standard ASCII.
    4. Replaces em-dashes and en-dashes with space-padded hyphens so words do not merge.
    5. Removes repeated SEC 10-K page headers and footers (e.g., 'Apple Inc. | 2025 Form 10-K | 21').
    6. Removes standalone page numbers (e.g., '21', 'Page 21').
    7. Removes isolated formatting symbols, bullets ('•'), and currency symbols on their own line.
    8. Removes standalone footnote markers on their own line (e.g., '(1)', '[a]').
    9. Removes isolated table-formatting artifacts (lines containing solely numbers or percentages).
    10. Normalizes internal multiple spaces, tabs, and line breaks into standard double newlines.
    11. Strictly preserves meaningful words, negative words ('not', 'no', 'never', 'without'),
        financial terminology, and readable sentence punctuation.

    Args:
        raw_text (str): Raw MD&A text loaded from extracted_text/<TICKER>/.

    Returns:
        tuple: (cleaned_text, orig_words, clean_words, words_removed, retained_pct)
    """
    if not raw_text or not raw_text.strip():
        return "", 0, 0, 0, 0.0

    orig_words = len(raw_text.split())

    # Step 1: Unescape HTML entities (&amp;, &nbsp;, &#160;, etc.) and remove stray HTML tags
    text = html.unescape(raw_text)
    text = re.sub(r"<[^>]+>", " ", text)

    # Step 2: Unicode normalization (NFKD)
    text = unicodedata.normalize("NFKD", text)

    # Step 3: Standardize apostrophes, quotes, and dashes
    # Convert curly single quotes and backticks to standard ASCII apostrophe '
    text = re.sub(r"[\u2018\u2019\u201a\u201b\u2032`]", "'", text)
    # Convert curly double quotes to standard ASCII quotation mark "
    text = re.sub(r"[\u201c\u201d\u201e\u201f\u2033«»]", '"', text)
    # Replace em-dashes and en-dashes with space-padded hyphen so words do not merge
    text = re.sub(r"[\u2013\u2014]", " - ", text)

    # Step 4: Line-by-line filtering of SEC artifacts, headers/footers, and isolated noise
    cleaned_lines = []
    lines = text.splitlines()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # Rule A: Remove SEC 10-K page headers and footers
        # Examples: "Apple Inc. | 2025 Form 10-K | 21", "Form 10-K | 45", "Item 7 | Page 12"
        if re.search(r"\|\s*\d{4}\s+Form\s+10-K\s*\|\s*\d+", stripped, re.IGNORECASE) or \
           re.search(r"Form\s+10-K\s*\|\s*\d+", stripped, re.IGNORECASE) or \
           re.match(r"^.*?\bForm\s+10-K\s*\|\s*\d+\s*$", stripped, re.IGNORECASE) or \
           re.match(r"^(?:Table of Contents|Index to Financial Statements)$", stripped, re.IGNORECASE):
            continue

        # Rule B: Remove standalone page numbers
        # Examples: "21", "Page 21", "p. 21", "Page 21 of 80"
        if re.match(r"^(?:Page\s+|p\.\s*)?\d{1,3}(?:\s+of\s+\d+)?$", stripped, re.IGNORECASE):
            continue

        # Rule C: Remove isolated formatting symbols and bullets
        # Examples: "•", "$", "%", "—", "-", "*", "_", "|", ":", "..."
        if re.match(r"^[\$\€\£\¥\%\•\*\_\—\–\|\-\:\.]+$", stripped):
            continue

        # Rule D: Remove isolated footnote markers on their own line
        # Examples: "(1)", "(2)", "(a)", "[1]"
        if re.match(r"^[\(\[]\s*[0-9a-zA-Z]{1,2}\s*[\)\]]$", stripped):
            continue

        # Rule E: Remove table rows consisting solely of isolated numbers, dollar amounts, or percentages
        # Examples: "178,353", "$178,353", "7%", "(4)%", "(8)", "36.8%"
        if re.match(r"^\(?\s*-?\s*\$?\s*\d{1,3}(?:,\d{3})*(?:\.\d+)?\s*%?\s*\)?$", stripped):
            continue

        # Rule F: Normalize horizontal whitespace within the line (convert tabs and multiple spaces to single space)
        cleaned_line = re.sub(r"[ \t]+", " ", stripped)

        # Rule G: Clean repetitive divider punctuation chains (e.g., "......", "------", "______")
        cleaned_line = re.sub(r"[\.]{4,}", "...", cleaned_line)
        cleaned_line = re.sub(r"[\-]{4,}", "---", cleaned_line)
        cleaned_line = re.sub(r"[\_]{4,}", "", cleaned_line).strip()

        if cleaned_line:
            cleaned_lines.append(cleaned_line)

    # Step 5: Rejoin into readable paragraphs with double newlines
    cleaned_text = "\n\n".join(cleaned_lines)
    cleaned_words = len(cleaned_text.split())
    words_removed = max(0, orig_words - cleaned_words)
    retained_pct = (cleaned_words / orig_words * 100.0) if orig_words > 0 else 0.0

    return cleaned_text, orig_words, cleaned_words, words_removed, retained_pct


def process_cleaned_mda_files(ticker, extracted_folder, cleaned_folder, filings_metadata=None):
    """
    Process extracted MD&A files for the current request, cleaning the text
    and saving to cleaned_text/<TICKER>/<TICKER>_<YEAR>_MDA_CLEAN.txt.

    Includes validation checks, aggressive cleaning safeguards, and error handling.

    Args:
        ticker (str): Company ticker symbol.
        extracted_folder (str): Directory containing raw extracted MD&A text files.
        cleaned_folder (str): Target directory for cleaned MD&A text files.
        filings_metadata (list, optional): Metadata from Step 3 for currently requested filings.

    Returns:
        list: Summary dictionary for each processed filing.
    """
    # Scope exclusively to the filings associated with the current user request
    if filings_metadata:
        target_years = [item["year"] for item in filings_metadata]
    else:
        # Fallback if metadata not supplied: discover from extracted_folder
        all_files = [f for f in os.listdir(extracted_folder) if f.startswith(f"{ticker}_") and f.endswith("_MDA.txt")]
        target_years = [f.replace(f"{ticker}_", "").replace("_MDA.txt", "") for f in sorted(all_files, reverse=True)]

    results = []
    print(f"\nProcessing MD&A text cleaning for {len(target_years)} filing(s)...")

    for year in target_years:
        source_filename = f"{ticker}_{year}_MDA.txt"
        clean_filename = f"{ticker}_{year}_MDA_CLEAN.txt"
        source_filepath = os.path.join(extracted_folder, source_filename)
        clean_filepath = os.path.join(cleaned_folder, clean_filename)

        # Validation Check 1: Ensure extracted source file exists
        if not os.path.isfile(source_filepath):
            print(f"  [ERROR] Source MD&A file '{source_filename}' not found.")
            results.append({
                "source_file": source_filename,
                "clean_file": clean_filename,
                "orig_words": 0,
                "clean_words": 0,
                "words_removed": 0,
                "retained_pct": 0.0,
                "save_location": clean_filepath,
                "status": "Missing source file",
                "warning": True
            })
            continue

        # Skip rewriting existing cleaned file unless regeneration is required
        if os.path.exists(clean_filepath):
            try:
                with open(source_filepath, "r", encoding="utf-8") as sf:
                    orig_words = len(sf.read().split())
                with open(clean_filepath, "r", encoding="utf-8") as cf:
                    clean_words = len(cf.read().split())
                words_removed = max(0, orig_words - clean_words)
                retained_pct = (clean_words / orig_words * 100.0) if orig_words > 0 else 0.0
                status_str = "Already cleaned (Skipped)"
                warning_flag = False

                if retained_pct < 70.0:
                    status_str = "Already cleaned (Review)"
                    warning_flag = True

                print(f"  [SKIPPED] {clean_filename} already exists.")
                results.append({
                    "source_file": source_filename,
                    "clean_file": clean_filename,
                    "orig_words": orig_words,
                    "clean_words": clean_words,
                    "words_removed": words_removed,
                    "retained_pct": retained_pct,
                    "save_location": clean_filepath,
                    "status": status_str,
                    "warning": warning_flag
                })
                continue
            except OSError:
                # If error reading existing file, proceed to regenerate
                pass

        # Validation Check 2: Read source file with error handling
        try:
            with open(source_filepath, "r", encoding="utf-8") as sf:
                raw_text = sf.read()
        except OSError as read_err:
            print(f"  [ERROR] Unable to read source file '{source_filename}': {read_err}")
            results.append({
                "source_file": source_filename,
                "clean_file": clean_filename,
                "orig_words": 0,
                "clean_words": 0,
                "words_removed": 0,
                "retained_pct": 0.0,
                "save_location": clean_filepath,
                "status": "Read error",
                "warning": True
            })
            continue

        # Validation Check 3: Check for empty source file
        if not raw_text.strip():
            print(f"  [ERROR] Source MD&A file '{source_filename}' is empty.")
            results.append({
                "source_file": source_filename,
                "clean_file": clean_filename,
                "orig_words": 0,
                "clean_words": 0,
                "words_removed": 0,
                "retained_pct": 0.0,
                "save_location": clean_filepath,
                "status": "Empty source file",
                "warning": True
            })
            continue

        # Execute text cleaning pipeline
        cleaned_text, orig_words, clean_words, words_removed, retained_pct = clean_mda_text(raw_text)

        # Validation Check 4: Check if cleaning produced an empty file
        if not cleaned_text.strip():
            print(f"  [ERROR] Cleaning produced an empty text file for '{source_filename}'.")
            results.append({
                "source_file": source_filename,
                "clean_file": clean_filename,
                "orig_words": orig_words,
                "clean_words": 0,
                "words_removed": orig_words,
                "retained_pct": 0.0,
                "save_location": clean_filepath,
                "status": "Cleaned file is empty",
                "warning": True
            })
            continue

        # Safeguard Check: Detect aggressive cleaning (>30% of original words removed)
        warning_flag = False
        if retained_pct < 70.0:
            warning_flag = True
            print(f"  [WARNING] {clean_filename} retained only {retained_pct:.1f}% of words (>30% removed). Manual review recommended.")
            status_str = "Cleaned (Review)"
        else:
            status_str = "Cleaned"

        # Save cleaned MD&A text to destination file
        try:
            with open(clean_filepath, "w", encoding="utf-8") as cf:
                cf.write(cleaned_text)
            print(f"  [CLEANED] Saved {clean_filename} ({clean_words:,d} words, {retained_pct:.1f}% retained)")
            results.append({
                "source_file": source_filename,
                "clean_file": clean_filename,
                "orig_words": orig_words,
                "clean_words": clean_words,
                "words_removed": words_removed,
                "retained_pct": retained_pct,
                "save_location": clean_filepath,
                "status": status_str,
                "warning": warning_flag
            })
        except OSError as write_err:
            print(f"  [ERROR] Could not save cleaned file '{clean_filename}': {write_err}")
            results.append({
                "source_file": source_filename,
                "clean_file": clean_filename,
                "orig_words": orig_words,
                "clean_words": clean_words,
                "words_removed": words_removed,
                "retained_pct": retained_pct,
                "save_location": clean_filepath,
                "status": "Save error",
                "warning": True
            })

    return results


def display_step6_summary(ticker, cleaned_folder, results):
    """
    Display a formatted confirmation table for Step 6 in the terminal.

    Args:
        ticker (str): Company ticker symbol.
        cleaned_folder (str): Directory where cleaned text files are stored.
        results (list): List of dictionaries containing file metrics from process_cleaned_mda_files.
    """
    print("\n" + "=" * 80)
    print("                    STEP 6: MD&A TEXT CLEANING SUMMARY")
    print("=" * 80)
    print(f"Company Ticker Symbol : {ticker}")
    print(f"Cleaned Directory Path: {cleaned_folder}")
    print("-" * 80)
    print(f"{'Source File':<24} {'Original Words':<16} {'Cleaned Words':<16} {'Retained':<12} {'Status'}")
    print("-" * 80)

    for item in results:
        src = item["source_file"]
        orig = f"{item['orig_words']:,d}"
        clean = f"{item['clean_words']:,d}"
        ret = f"{item['retained_pct']:.1f}%"
        status = item["status"]
        print(f"{src:<24} {orig:<16} {clean:<16} {ret:<12} {status}")

    print("=" * 80)
    print("Step 6 complete! Cleaned MD&A text is prepared for sentiment analysis.\n")


# ==============================================================================
# STEP 7: LOUGHRAN-MCDONALD SENTIMENT ANALYSIS FUNCTIONS
# ==============================================================================

def setup_dictionaries_directory():
    """
    Ensure the 'dictionaries' folder exists in the project directory.

    Returns:
        str: Absolute path to the dictionaries folder.
    """
    dict_dir = str(PROJECT_ROOT / "dictionaries")
    try:
        os.makedirs(dict_dir, exist_ok=True)
        return dict_dir
    except OSError as err:
        print(f"\nError: Could not create dictionaries directory '{dict_dir}': {err}")
        return dict_dir


def find_dictionary_file(dictionaries_dir):
    """
    Search dictionaries_dir for a Loughran-McDonald dictionary file (.csv or .xlsx).

    Args:
        dictionaries_dir (str): Path to dictionaries folder.

    Returns:
        tuple: (dict_filepath, dict_filename) or (None, None) if not found.
    """
    if not os.path.isdir(dictionaries_dir):
        return None, None

    valid_exts = (".csv", ".xlsx")
    candidates = [
        f for f in os.listdir(dictionaries_dir)
        if any(f.lower().endswith(ext) for ext in valid_exts) and not f.startswith("~$")
    ]

    if not candidates:
        return None, None

    # Prioritize files containing 'loughran', 'mcdonald', 'master', or 'sentiment'
    lm_candidates = [
        f for f in candidates
        if any(term in f.lower() for term in ["loughran", "mcdonald", "master", "sentiment"])
    ]
    chosen = sorted(lm_candidates)[0] if lm_candidates else sorted(candidates)[0]

    if len(candidates) > 1:
        print(f"  Note: Multiple dictionary files detected in '{dictionaries_dir}'. Using '{chosen}'.")

    return os.path.join(dictionaries_dir, chosen), chosen


def display_missing_dictionary_notice(dictionaries_dir):
    """
    Display a clear terminal notification when the Loughran-McDonald dictionary
    file has not yet been placed in the dictionaries folder.
    """
    print("\n" + "=" * 80)
    print("             STEP 7: LOUGHRAN-MCDONALD DICTIONARY REQUIRED")
    print("=" * 80)
    print("Notice: Loughran-McDonald financial sentiment dictionary file not found!")
    print("\nPlease place your Loughran-McDonald dictionary file (.csv or .xlsx) in the folder:")
    print(f"  {dictionaries_dir}")
    print("\nSupported formats:")
    print("  - CSV file   (e.g., Loughran-McDonald_MasterDictionary_1993-2023.csv)")
    print("  - Excel file (e.g., Loughran-McDonald_MasterDictionary.xlsx)")
    print("\nOnce placed in the 'dictionaries' folder, simply re-run:")
    print("  python3 main.py")
    print("=" * 80 + "\n")


def is_category_member(val):
    """
    Check if a dictionary cell value indicates category membership.
    Handles ND-SRAF year integers (e.g., 2009, 1993), 1/0, True/False, 'Yes', etc.
    """
    if val is None:
        return False
    if isinstance(val, bool):
        return val
    if isinstance(val, (int, float)):
        return val != 0 and not math.isnan(val)
    val_str = str(val).strip().lower()
    if val_str in ("", "0", "0.0", "false", "none", "nan", "null", "no", "n"):
        return False
    try:
        num = float(val_str)
        return num != 0
    except ValueError:
        return True


def _parse_dictionary_rows(dict_rows, headers, dict_filename, dict_path):
    """
    Internal helper to extract words and available sentiment categories from
    either CSV or Excel row dictionaries.
    """
    CATEGORY_MAPPING = {
        "positive": "Positive",
        "negative": "Negative",
        "uncertainty": "Uncertainty",
        "litigious": "Litigious",
        "strong modal": "Strong Modal",
        "strong_modal": "Strong Modal",
        "strongmodal": "Strong Modal",
        "weak modal": "Weak Modal",
        "weak_modal": "Weak Modal",
        "weakmodal": "Weak Modal",
        "constraining": "Constraining",
        "complexity": "Complexity",
    }

    # Identify word column
    word_col = None
    for h in headers:
        clean_h = re.sub(r"[\s_]+", " ", str(h).strip().lower())
        if clean_h in ("word", "words", "term", "token"):
            word_col = h
            break

    if not word_col and headers:
        word_col = headers[0]

    # Identify available sentiment category columns
    matched_categories = {}
    for h in headers:
        clean_h = re.sub(r"[\s_]+", " ", str(h).strip().lower())
        if clean_h in CATEGORY_MAPPING:
            cat_name = CATEGORY_MAPPING[clean_h]
            matched_categories[cat_name] = h

    if not matched_categories:
        print(f"\nError: No recognizable Loughran-McDonald sentiment category columns found in '{dict_filename}'.")
        return None

    # Canonical presentation order
    canonical_order = [
        "Positive", "Negative", "Uncertainty", "Litigious",
        "Strong Modal", "Weak Modal", "Constraining", "Complexity"
    ]
    available_categories = [c for c in canonical_order if c in matched_categories]
    for c in matched_categories:
        if c not in available_categories:
            available_categories.append(c)

    category_sets = {cat: set() for cat in available_categories}
    all_dict_words = set()

    for row in dict_rows:
        raw_word = row.get(word_col)
        if raw_word is None:
            continue
        word = str(raw_word).strip().upper()
        if not word or word == "NAN":
            continue
        all_dict_words.add(word)

        for cat in available_categories:
            col_name = matched_categories[cat]
            val = row.get(col_name)
            if is_category_member(val):
                category_sets[cat].add(word)

    return {
        "filename": dict_filename,
        "filepath": dict_path,
        "categories": available_categories,
        "category_sets": category_sets,
        "total_words": len(all_dict_words)
    }


def _load_csv_dictionary(dict_path, dict_filename):
    """
    Load Loughran-McDonald dictionary from a CSV file.
    """
    encoding_list = ["utf-8-sig", "utf-8", "latin-1", "cp1252"]
    reader_rows = None

    for enc in encoding_list:
        try:
            with open(dict_path, "r", encoding=enc) as f:
                reader = csv.DictReader(f)
                reader_rows = list(reader)
                headers = reader.fieldnames or []
                break
        except (UnicodeDecodeError, OSError):
            continue

    if reader_rows is None or not headers:
        print(f"\nError: Could not read CSV dictionary from '{dict_path}'.")
        return None

    return _parse_dictionary_rows(reader_rows, headers, dict_filename, dict_path)


def _load_excel_dictionary(dict_path, dict_filename):
    """
    Load Loughran-McDonald dictionary from an Excel file (.xlsx) using openpyxl.
    """
    try:
        wb = openpyxl.load_workbook(dict_path, read_only=True, data_only=True)
        sheet = wb.active
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if not header_row:
            print(f"\nError: Excel dictionary '{dict_filename}' is empty.")
            return None

        headers = [str(c).strip() if c is not None else "" for c in header_row]
        dict_rows = []
        for r in rows:
            if r is None or not any(r):
                continue
            row_dict = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
            dict_rows.append(row_dict)

        return _parse_dictionary_rows(dict_rows, headers, dict_filename, dict_path)
    except Exception as err:
        print(f"\nError reading Excel dictionary '{dict_filename}': {err}")
        return None


def load_loughran_mcdonald_dictionary(dict_path):
    """
    Load and index the Loughran-McDonald sentiment dictionary from CSV (.csv) or Excel (.xlsx).

    Args:
        dict_path (str): Full path to dictionary file.

    Returns:
        dict: Parsed dictionary structure with word sets per category, or None on error.
    """
    dict_filename = os.path.basename(dict_path)
    file_ext = os.path.splitext(dict_filename)[1].lower()

    if file_ext == ".csv":
        return _load_csv_dictionary(dict_path, dict_filename)
    elif file_ext == ".xlsx":
        return _load_excel_dictionary(dict_path, dict_filename)
    elif file_ext == ".xls":
        print(f"\nError: Legacy Excel format '{file_ext}' is not supported by openpyxl. Please save or convert your dictionary as .xlsx or .csv.")
        return None
    else:
        print(f"\nError: Unsupported dictionary format '{file_ext}'. Please provide a .csv or .xlsx file.")
        return None


def tokenize_mda_words(text):
    """
    Tokenize cleaned MD&A text into individual uppercase words for dictionary matching.
    Handles apostrophes consistently (e.g., possessive 'S suffix is removed so
    base words match dictionary entries).

    No stemming, lemmatization, or fuzzy matching is applied.

    Args:
        text (str): Cleaned MD&A text.

    Returns:
        list: List of uppercase alphabetic word tokens.
    """
    raw_tokens = re.findall(r"\b[A-Za-z]+(?:'[A-Za-z]+)?\b", text)
    tokens = []
    for t in raw_tokens:
        clean_t = t.upper().strip("'")
        if clean_t.endswith("'S"):
            clean_t = clean_t[:-2]
        if clean_t:
            tokens.append(clean_t)
    return tokens


def analyze_mda_sentiment(cleaned_folder, ticker, filings_metadata, dict_info):
    """
    Score the cleaned MD&A text files against the Loughran-McDonald dictionary.
    Only processes filings associated with the current user request.

    Calculates:
    - Total Words
    - Word counts for each available sentiment category
    - Positive-to-Negative (P/N) Ratio (with zero-division protection)
    - Normalized percentage share for each sentiment category

    Args:
        cleaned_folder (str): Folder containing cleaned MD&A text files.
        ticker (str): Company ticker symbol.
        filings_metadata (list): Metadata of filings for the current request.
        dict_info (dict): Loaded Loughran-McDonald dictionary metadata and word sets.

    Returns:
        list: Structured sentiment results per fiscal year sorted newest to oldest.
    """
    if filings_metadata:
        target_years = [item["year"] for item in filings_metadata]
    else:
        clean_files = [f for f in os.listdir(cleaned_folder) if f.startswith(f"{ticker}_") and f.endswith("_MDA_CLEAN.txt")]
        target_years = [f.replace(f"{ticker}_", "").replace("_MDA_CLEAN.txt", "") for f in sorted(clean_files, reverse=True)]

    # Files are named by filing year; results are labeled by the fiscal year reported
    fiscal_year_map = {
        str(item["year"]): str(item.get("fiscal_year") or item["year"])
        for item in (filings_metadata or [])
    }

    results = []
    available_cats = dict_info["categories"]
    category_sets = dict_info["category_sets"]

    print(f"\nAnalyzing financial sentiment across {len(target_years)} filing(s)...")

    for year in target_years:
        clean_filename = f"{ticker}_{year}_MDA_CLEAN.txt"
        clean_filepath = os.path.join(cleaned_folder, clean_filename)

        if not os.path.isfile(clean_filepath):
            print(f"  [ERROR] Cleaned file '{clean_filename}' not found. Skipping.")
            continue

        try:
            with open(clean_filepath, "r", encoding="utf-8") as f:
                content = f.read()
        except OSError as read_err:
            print(f"  [ERROR] Unable to read '{clean_filename}': {read_err}")
            continue

        tokens = tokenize_mda_words(content)
        total_words = len(tokens)

        # Sanity Check 1: Empty text
        if total_words == 0:
            print(f"  [WARNING] Total word count is zero for '{clean_filename}'.")
            counts = {cat: 0 for cat in available_cats}
            pn_ratio = None
            pcts = {cat: 0.0 for cat in available_cats}
            status_str = "Zero Words"
        else:
            # Calculate counts for each category
            counts = {}
            for cat in available_cats:
                cat_words = category_sets[cat]
                counts[cat] = sum(1 for t in tokens if t in cat_words)

            # Sanity Check 2: Check if any category was matched
            total_sentiment_matches = sum(counts.values())
            if total_sentiment_matches == 0:
                print(f"  [WARNING] No words matched any Loughran-McDonald category in '{clean_filename}'.")

            pos_count = counts.get("Positive", 0)
            neg_count = counts.get("Negative", 0)

            # P/N Ratio calculation with safe zero-division handling
            if neg_count > 0:
                pn_ratio = round(pos_count / neg_count, 3)
            else:
                pn_ratio = None  # Documented representation for zero negative words

            # Category percentages of total words
            pcts = {}
            for cat in available_cats:
                pcts[cat] = (counts[cat] / total_words) if total_words > 0 else 0.0

            status_str = "Analyzed"

        results.append({
            "year": fiscal_year_map.get(str(year), str(year)),
            "source_file": clean_filename,
            "total_words": total_words,
            "counts": counts,
            "pn_ratio": pn_ratio,
            "percentages": pcts,
            "status": status_str
        })

    # Sort results by fiscal year from newest to oldest
    results.sort(key=lambda x: str(x["year"]), reverse=True)
    return results


def add_sentiment_worksheet_to_workbook(company_name, ticker, sentiment_results, dict_info, output_filepath):
    """
    Add or refresh the 'Sentiment Analysis' worksheet in the existing Excel workbook.
    Preserves all existing financial statement tabs intact.

    Args:
        company_name (str): Company name.
        ticker (str): Company ticker symbol.
        sentiment_results (list): Sentiment metrics list per fiscal year.
        dict_info (dict): Loughran-McDonald dictionary information.
        output_filepath (str): Path to existing Excel file.

    Returns:
        bool: True if updated and saved successfully, False otherwise.
    """
    if not os.path.isfile(output_filepath):
        print(f"\nError: Excel workbook '{output_filepath}' not found.")
        return False

    try:
        workbook = openpyxl.load_workbook(output_filepath)
    except Exception as err:
        print(f"\nError opening workbook '{output_filepath}': {err}")
        return False

    # Safely refresh or replace existing Sentiment Analysis tab
    sheet_title = "Sentiment Analysis"
    if sheet_title in workbook.sheetnames:
        del workbook[sheet_title]

    worksheet = workbook.create_sheet(title=sheet_title)

    # Professional styles matching Step 5 tabs
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="595959")
    regular_font = Font(name="Calibri", size=10, color="000000")

    # Row 1: Title Header
    worksheet["A1"] = f"{company_name.upper()} - LOUGHRAN-MCDONALD FINANCIAL SENTIMENT ANALYSIS"
    worksheet["A1"].font = title_font

    # Row 2: Subtitle Header
    dict_fname = dict_info.get("filename", "Loughran-McDonald Dictionary")
    worksheet["A2"] = f"Dictionary: {dict_fname} | Source: Item 7 MD&A Cleaned Text"
    worksheet["A2"].font = subtitle_font

    # Row 4: Column Headers
    available_cats = dict_info["categories"]
    headers = ["Fiscal Year", "Total Words"] + available_cats + ["P/N Ratio"] + [f"{cat} %" for cat in available_cats]

    for col_idx, h in enumerate(headers, 1):
        cell = worksheet.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "right", vertical="center")

    # Rows 5+: Data rows
    for r_idx, res in enumerate(sentiment_results, 5):
        year_val = res["year"]
        tot_words = res["total_words"]
        counts = res["counts"]
        pn_val = res["pn_ratio"]
        pcts = res["percentages"]

        col_pos = 1

        # Col 1: Fiscal Year
        c = worksheet.cell(row=r_idx, column=col_pos, value=year_val)
        c.font = regular_font
        c.alignment = Alignment(horizontal="center")
        col_pos += 1

        # Col 2: Total Words
        c = worksheet.cell(row=r_idx, column=col_pos, value=tot_words)
        c.font = regular_font
        c.alignment = Alignment(horizontal="right")
        c.number_format = "#,##0"
        col_pos += 1

        # Category word counts
        for cat in available_cats:
            val = counts.get(cat, 0)
            c = worksheet.cell(row=r_idx, column=col_pos, value=val)
            c.font = regular_font
            c.alignment = Alignment(horizontal="right")
            c.number_format = "#,##0"
            col_pos += 1

        # P/N Ratio
        if pn_val is not None:
            c = worksheet.cell(row=r_idx, column=col_pos, value=pn_val)
            c.number_format = "0.00"
        else:
            c = worksheet.cell(row=r_idx, column=col_pos, value="N/A")
        c.font = regular_font
        c.alignment = Alignment(horizontal="right")
        col_pos += 1

        # Category percentages
        for cat in available_cats:
            pct_val = pcts.get(cat, 0.0)
            c = worksheet.cell(row=r_idx, column=col_pos, value=pct_val)
            c.font = regular_font
            c.alignment = Alignment(horizontal="right")
            c.number_format = "0.00%"
            col_pos += 1

    # Freeze header rows at row 5
    worksheet.freeze_panes = "B5"
    worksheet.views.sheetView[0].showGridLines = True

    # Auto-adjust column widths
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        worksheet.column_dimensions[col_letter].width = max(14, len(headers[col_idx - 1]) + 4)

    try:
        workbook.save(output_filepath)
        return True
    except OSError as save_err:
        print(f"\nError saving Excel workbook '{output_filepath}': {save_err}")
        return False


def display_step7_summary(ticker, dict_info, workbook_path, sentiment_results):
    """
    Display a formatted confirmation table for Step 7 in the terminal.
    """
    print("\n" + "=" * 80)
    print("               STEP 7: LOUGHRAN-MCDONALD SENTIMENT ANALYSIS SUMMARY")
    print("=" * 80)
    print(f"Company Ticker Symbol : {ticker}")
    print(f"Dictionary File Used  : {dict_info['filename']}")
    print(f"Dictionary Words Loaded: {dict_info['total_words']:,d}")
    print(f"Excel Workbook Path   : {workbook_path}")
    print("Worksheet Tab Created : Sentiment Analysis")
    print("-" * 80)

    # Dynamic header based on available categories
    available_cats = dict_info["categories"]
    header_str = f"{'Fiscal Year':<12} {'Total Words':<13} "
    for cat in available_cats[:4]:  # Show primary categories in terminal table
        header_str += f"{cat:<12} "
    header_str += f"{'P/N Ratio':<10} {'Status':<10}"

    print(header_str)
    print("-" * 80)

    for res in sentiment_results:
        yr = res["year"]
        tot = f"{res['total_words']:,d}"
        row_str = f"{yr:<12} {tot:<13} "
        for cat in available_cats[:4]:
            c_val = f"{res['counts'].get(cat, 0):,d}"
            row_str += f"{c_val:<12} "
        pn_str = f"{res['pn_ratio']:.2f}" if res["pn_ratio"] is not None else "N/A"
        row_str += f"{pn_str:<10} {res['status']:<10}"
        print(row_str)

    print("=" * 80)
    print("Step 7 complete! Sentiment analysis worksheet added to Excel workbook.\n")


# ==============================================================================
# STEP 8: HISTORICAL STOCK PRICE RETRIEVAL FUNCTIONS
# ==============================================================================

def retrieve_stock_prices(ticker, filings_metadata):
    """
    Retrieve historical split-adjusted closing stock prices associated with each
    company's actual balance sheet date (report_date).
    Uses yfinance to query Yahoo Finance without requiring an API key.

    If the balance sheet date falls on a non-trading day (weekend or market holiday),
    the function identifies and uses the most recent previous trading day.
    Future dates after the balance sheet date are never used.

    Args:
        ticker (str): Company ticker symbol.
        filings_metadata (list): List of filing metadata dicts containing 'year' and 'report_date'.

    Returns:
        list of dict: Stock price records sorted newest fiscal year to oldest, or None on error.
    """
    try:
        import yfinance as yf
    except ImportError:
        print("\n" + "!" * 80)
        print("ERROR: The 'yfinance' library is required to retrieve stock prices.")
        print("Please install it in your terminal using:")
        print("    python3 -m pip install yfinance")
        print("!" * 80 + "\n")
        return None

    if not filings_metadata:
        print(f"\nWarning: No filing metadata provided for stock price retrieval.")
        return None

    print(f"\nRetrieving historical stock prices for ticker '{ticker}' from Yahoo Finance...")

    # Parse all target balance sheet dates
    target_periods = []
    for item in filings_metadata:
        # Label by fiscal year (balance sheet year), not the SEC filing year
        yr = str(item.get("fiscal_year") or item.get("year", ""))
        # Use actual balance sheet date from SEC EDGAR report_date
        bs_date_str = item.get("report_date") or item.get("filing_date")
        try:
            bs_date = datetime.strptime(bs_date_str, "%Y-%m-%d").date()
        except (ValueError, TypeError):
            print(f"  [WARNING] Invalid balance sheet date '{bs_date_str}' for fiscal year {yr}.")
            continue
        target_periods.append({
            "year": yr,
            "bs_date": bs_date,
            "bs_date_str": bs_date.strftime("%Y-%m-%d")
        })

    if not target_periods:
        print("Error: Could not determine valid balance sheet dates.")
        return None

    # Calculate date window for yfinance download:
    # From 30 days before the oldest balance sheet date up to 7 days after the newest
    all_dates = [tp["bs_date"] for tp in target_periods]
    earliest_date = min(all_dates) - timedelta(days=30)
    latest_date = max(all_dates) + timedelta(days=7)

    start_str = earliest_date.strftime("%Y-%m-%d")
    end_str = latest_date.strftime("%Y-%m-%d")

    try:
        ticker_obj = yf.Ticker(ticker.upper())
        # auto_adjust=False ensures standard Close is split-adjusted
        hist = ticker_obj.history(start=start_str, end=end_str, auto_adjust=False)
    except Exception as err:
        print(f"\nError connecting to Yahoo Finance for ticker '{ticker}': {err}")
        return None

    if hist is None or hist.empty:
        print(f"\nError: No historical stock price data returned for ticker '{ticker}'.")
        return None

    # Map trading dates in history
    results = []
    for tp in target_periods:
        yr = tp["year"]
        bs_date = tp["bs_date"]
        bs_date_str = tp["bs_date_str"]

        # Filter trading days strictly ON or BEFORE the balance sheet date
        valid_trading = hist[hist.index.date <= bs_date]

        if valid_trading.empty:
            print(f"  [WARNING] No historical trading days found on or before {bs_date_str} for {yr}.")
            results.append({
                "year": yr,
                "bs_date": bs_date_str,
                "trading_date": "N/A",
                "price": None,
                "source": "Yahoo Finance",
                "adjustment": "No data",
                "status": "Missing Data"
            })
            continue

        # Last trading day on or prior to the balance sheet date
        last_trade_dt = valid_trading.index[-1].date()
        last_row = valid_trading.iloc[-1]

        try:
            raw_price = float(last_row["Close"])
            stock_price = round(raw_price, 2)
        except (ValueError, KeyError, TypeError):
            stock_price = None

        if last_trade_dt == bs_date:
            adjustment_desc = "Exact date"
        else:
            adjustment_desc = "Previous trading day"

        if stock_price is None or stock_price <= 0:
            status_desc = "Invalid Price"
            print(f"  [WARNING] Non-positive or invalid stock price for {yr}: {stock_price}")
        else:
            status_desc = "Retrieved"

        results.append({
            "year": yr,
            "bs_date": bs_date_str,
            "trading_date": last_trade_dt.strftime("%Y-%m-%d"),
            "price": stock_price,
            "source": "Yahoo Finance",
            "adjustment": adjustment_desc,
            "status": status_desc
        })

    # Sort results newest fiscal year to oldest
    results.sort(key=lambda x: str(x["year"]), reverse=True)

    # Sanity check: confirm number of retrieved prices matches requested count
    if len(results) != len(filings_metadata):
        print(f"  [NOTICE] Retrieved stock prices for {len(results)} period(s) "
              f"(requested: {len(filings_metadata)}).")

    return results


def add_stock_prices_worksheet_to_workbook(company_name, ticker, stock_results, output_filepath):
    """
    Add or refresh the 'Stock Prices' worksheet in the existing Excel workbook.
    Preserves all existing worksheets (Income Statement, Balance Sheet,
    Cash Flows, Sentiment Analysis).

    Args:
        company_name (str): Company name.
        ticker (str): Stock ticker symbol.
        stock_results (list): List of stock price records.
        output_filepath (str): Path to Excel workbook.

    Returns:
        bool: True if created and saved successfully, False otherwise.
    """
    if not os.path.isfile(output_filepath):
        print(f"\nError: Excel workbook '{output_filepath}' not found.")
        return False

    try:
        workbook = openpyxl.load_workbook(output_filepath)
    except Exception as err:
        print(f"\nError opening workbook '{output_filepath}': {err}")
        return False

    # Safely refresh or replace existing Stock Prices tab
    sheet_title = "Stock Prices"
    if sheet_title in workbook.sheetnames:
        del workbook[sheet_title]

    worksheet = workbook.create_sheet(title=sheet_title)

    # Professional styles matching prior tabs
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="595959")
    regular_font = Font(name="Calibri", size=10, color="000000")

    # Row 1: Title Header
    worksheet["A1"] = f"{company_name.upper()} - HISTORICAL STOCK PRICE ANALYSIS"
    worksheet["A1"].font = title_font

    # Row 2: Subtitle Header
    worksheet["A2"] = "Price Type: Split-Adjusted Closing Price (Close) | Source: Yahoo Finance (yfinance)"
    worksheet["A2"].font = subtitle_font

    # Row 4: Column Headers
    headers = [
        "Fiscal Year",
        "Balance Sheet Date",
        "Trading Date Used",
        "Stock Price",
        "Source",
        "Date Adjustment"
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = worksheet.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        align_pos = "center" if col_idx in (1, 2, 3, 5) else ("right" if col_idx == 4 else "left")
        cell.alignment = Alignment(horizontal=align_pos, vertical="center")

    # Rows 5+: Data rows
    for r_idx, row in enumerate(stock_results, 5):
        # Col 1: Fiscal Year
        c1 = worksheet.cell(row=r_idx, column=1, value=row["year"])
        c1.font = regular_font
        c1.alignment = Alignment(horizontal="center")

        # Col 2: Balance Sheet Date
        c2 = worksheet.cell(row=r_idx, column=2, value=row["bs_date"])
        c2.font = regular_font
        c2.alignment = Alignment(horizontal="center")

        # Col 3: Trading Date Used
        c3 = worksheet.cell(row=r_idx, column=3, value=row["trading_date"])
        c3.font = regular_font
        c3.alignment = Alignment(horizontal="center")

        # Col 4: Stock Price (Numeric float formatted as currency)
        price_val = row["price"]
        if price_val is not None:
            c4 = worksheet.cell(row=r_idx, column=4, value=price_val)
            c4.number_format = "$#,##0.00"
        else:
            c4 = worksheet.cell(row=r_idx, column=4, value="N/A")
        c4.font = regular_font
        c4.alignment = Alignment(horizontal="right")

        # Col 5: Source
        c5 = worksheet.cell(row=r_idx, column=5, value=row["source"])
        c5.font = regular_font
        c5.alignment = Alignment(horizontal="center")

        # Col 6: Date Adjustment
        c6 = worksheet.cell(row=r_idx, column=6, value=row["adjustment"])
        c6.font = regular_font
        c6.alignment = Alignment(horizontal="left")

    # Freeze header rows at row 5
    worksheet.freeze_panes = "B5"
    worksheet.views.sheetView[0].showGridLines = True

    # Auto-adjust column widths
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        worksheet.column_dimensions[col_letter].width = max(16, len(headers[col_idx - 1]) + 4)

    try:
        workbook.save(output_filepath)
        return True
    except OSError as save_err:
        print(f"\nError saving Excel workbook '{output_filepath}': {save_err}")
        return False


def display_step8_summary(ticker, workbook_path, stock_results):
    """
    Display a formatted confirmation table for Step 8 in the terminal.
    """
    print("\n" + "=" * 80)
    print("                     STEP 8: STOCK PRICE RETRIEVAL SUMMARY")
    print("=" * 80)
    print(f"Company Ticker Symbol : {ticker}")
    print("Data Source           : Yahoo Finance (yfinance)")
    print("Price Methodology     : Split-Adjusted Closing Price (Close)")
    print(f"Excel Workbook Path   : {workbook_path}")
    print("Worksheet Tab Created : Stock Prices")
    print("-" * 80)
    print(f"{'Fiscal Year':<12} {'Balance Sheet Date':<20} {'Trading Date Used':<20} {'Stock Price':<14} {'Status'}")
    print("-" * 80)

    for item in stock_results:
        yr = item["year"]
        bs = item["bs_date"]
        td = item["trading_date"]
        pr = f"${item['price']:,.2f}" if item["price"] is not None else "N/A"
        st = item["status"]
        print(f"{yr:<12} {bs:<20} {td:<20} {pr:<14} {st}")

    print("=" * 80)
    print("Step 8 complete! Historical stock prices added to Excel workbook.\n")


# ==============================================================================
# STEP 9: FINAL FINANCIAL RATIOS, SENTIMENT & STOCK PRICE VISUALIZATION
# ==============================================================================

def find_statement_account_data(statement, patterns, exclude_terms=None):
    """
    Search a financial statement for the most appropriate line item matching
    the desired financial metric patterns, ranking exact matches over substrings
    and checking negative exclusions (e.g., excluding 'per share' or 'equity').

    Args:
        statement (dict): Statement dictionary with 'items' list and 'values' dict.
        patterns (list): Desired patterns in descending order of priority.
        exclude_terms (list, optional): Substrings that disqualify an item.

    Returns:
        tuple: (matched_label, values_dict, item_key) or (None, {}, None)
    """
    if not statement:
        return None, {}, None

    if exclude_terms is None:
        exclude_terms = ["per share", "basic", "diluted", "share:"]

    items_list = statement.get("items", [])
    values_dict = statement.get("values", {})

    candidates = []

    for item in items_list:
        key, label, is_header = item
        if is_header:
            continue

        norm_label = re.sub(r"[\s_]+", " ", label.strip().lower())
        norm_clean = re.sub(r"[^\w\s]", "", norm_label)

        # Check exclusions
        if any(ex in norm_label or ex in norm_clean for ex in exclude_terms):
            continue

        vals = values_dict.get(key, {})
        has_non_empty_vals = any(v is not None for v in vals.values())

        for priority, pattern in enumerate(patterns):
            norm_pattern = re.sub(r"[\s_]+", " ", pattern.strip().lower())
            clean_pattern = re.sub(r"[^\w\s]", "", norm_pattern)

            # Match criteria: exact match gets highest rank (0), substring gets (1)
            if norm_label == norm_pattern or norm_clean == clean_pattern:
                candidates.append((priority, 0, not has_non_empty_vals, key, label, vals))
                break
            elif clean_pattern in norm_clean or norm_pattern in norm_label:
                candidates.append((priority, 1, not has_non_empty_vals, key, label, vals))
                break

    if not candidates:
        return None, {}, None

    # Sort candidates by:
    # 1. Pattern priority (lower index in patterns list is better)
    # 2. Exact match (0) vs substring (1)
    # 3. Has non-empty values (False comes before True)
    candidates.sort(key=lambda x: (x[0], x[1], x[2]))
    best = candidates[0]
    return best[4], best[5], best[3]


def calculate_financial_ratios(statement_data, years):
    """
    Compute the four required financial ratios for each requested fiscal year:
      1. Liquidity:     Current Ratio = Current Assets / Current Liabilities
      2. Solvency:      Debt-to-Assets Ratio = Total Liabilities / Total Assets
      3. Efficiency:    Total Asset Turnover = Revenue / Average Total Assets
                        where Average Total Assets = (Beginning Assets + Ending Assets) / 2
      4. Profitability: Net Profit Margin = Net Income / Revenue

    Handles prior-year Total Assets retrieval from comparative balance sheet data
    for the oldest displayed year, division-by-zero protection, and missing accounts.

    Args:
        statement_data (dict): Dict of extracted financial statements.
        years (list): List of fiscal years to calculate ratios for.

    Returns:
        tuple: (ratios_by_year, raw_accounts_info)
    """
    is_stmt = statement_data.get("Income Statement", {})
    bs_stmt = statement_data.get("Balance Sheet", {})

    # Match the 6 underlying accounts
    rev_label, rev_vals, _ = find_statement_account_data(
        is_stmt,
        patterns=[
            "total net sales", "total revenue", "total revenues",
            "net sales", "revenue", "revenues", "operating revenues"
        ],
        exclude_terms=["cost", "costs", "deferred", "per share", "basic", "diluted"]
    )
    ni_label, ni_vals, _ = find_statement_account_data(
        is_stmt,
        patterns=[
            "net income", "net earnings", "net income loss",
            "net income/loss", "net profit"
        ],
        exclude_terms=[
            "per share", "basic", "diluted", "comprehensive",
            "operating income", "operating loss", "before", "loss per",
            "noncontrolling", "non-controlling"
        ]
    )
    # Balance sheet totals use only 'total ...' style labels so a component line
    # such as 'Other assets' or 'Accrued liabilities' can never be picked by mistake
    ca_label, ca_vals, _ = find_statement_account_data(
        bs_stmt,
        patterns=["total current assets", "current assets"],
        exclude_terms=["noncurrent", "non-current", "other current assets", "per share"]
    )
    cl_label, cl_vals, _ = find_statement_account_data(
        bs_stmt,
        patterns=["total current liabilities", "current liabilities"],
        exclude_terms=["noncurrent", "non-current", "other current liabilities", "per share"]
    )
    tl_label, tl_vals, _ = find_statement_account_data(
        bs_stmt,
        patterns=["total liabilities"],
        exclude_terms=["current liabilities", "equity", "shareholder", "stockholder", "per share"]
    )
    ta_label, ta_vals, _ = find_statement_account_data(
        bs_stmt,
        patterns=["total assets"],
        exclude_terms=["current assets", "noncurrent", "non-current", "equity", "liabilities and", "per share"]
    )

    # Some balance sheets have no 'Total liabilities' line. Derive it as
    # Total liabilities and equity - Total equity, and say so in the output.
    tle_label, tle_vals, _ = find_statement_account_data(
        bs_stmt,
        patterns=[
            "total liabilities and stockholders equity",
            "total liabilities and shareholders equity",
            "total liabilities and equity"
        ],
        exclude_terms=["per share"]
    )
    te_label, te_vals, _ = find_statement_account_data(
        bs_stmt,
        patterns=["total stockholders equity", "total shareholders equity", "total equity"],
        exclude_terms=["liabilities", "per share"]
    )
    tl_vals = dict(tl_vals)
    derived_tl_years = []
    if tle_label and te_label:
        for yr, tle in tle_vals.items():
            if tl_vals.get(yr) is None and te_vals.get(yr) is not None:
                tl_vals[yr] = tle - te_vals[yr]
                derived_tl_years.append(yr)
        if derived_tl_years and not tl_label:
            tl_label = f"{tle_label} - {te_label} (derived)"

    # Sanity checks for account detection
    missing_accounts = []
    if not rev_label:
        missing_accounts.append("Revenue / Net Sales")
    if not ni_label:
        missing_accounts.append("Net Income")
    if not ca_label:
        missing_accounts.append("Current Assets")
    if not cl_label:
        missing_accounts.append("Current Liabilities")
    if not tl_label:
        missing_accounts.append("Total Liabilities")
    if not ta_label:
        missing_accounts.append("Total Assets")

    if missing_accounts:
        print(f"  [WARNING] The following statement accounts could not be automatically identified: {', '.join(missing_accounts)}")

    ratios_by_year = {}
    underlying_values_by_year = {}

    def safe_divide(numerator, denominator, ratio_name, yr):
        # Missing inputs are reported per year below; a zero denominator is reported here
        if numerator is None or denominator is None:
            return None
        if denominator == 0:
            print(f"  [WARNING] FY {yr}: {ratio_name} not calculated (division by zero).")
            return None
        return numerator / denominator

    for yr in years:
        rev = rev_vals.get(yr)
        ni = ni_vals.get(yr)
        ca = ca_vals.get(yr)
        cl = cl_vals.get(yr)
        tl = tl_vals.get(yr)
        ta = ta_vals.get(yr)

        # Report every missing input for this year rather than substituting zero
        for acct_name, acct_val in [
            ("Current Assets", ca), ("Current Liabilities", cl),
            ("Total Liabilities", tl), ("Total Assets", ta),
            ("Revenue", rev), ("Net Income", ni)
        ]:
            if acct_val is None:
                print(f"  [WARNING] FY {yr}: {acct_name} missing - dependent ratio(s) shown as N/A.")
        if yr in derived_tl_years:
            print(f"  [NOTE] FY {yr}: Total Liabilities derived as {tle_label} minus {te_label}.")

        # 1. Liquidity: Current Ratio = Current Assets / Current Liabilities
        current_ratio = safe_divide(ca, cl, "Current Ratio", yr)

        # 2. Solvency: Debt-to-Assets Ratio = Total Liabilities / Total Assets
        debt_to_assets = safe_divide(tl, ta, "Debt-to-Assets Ratio", yr)

        # 3. Efficiency: Total Asset Turnover = Revenue / Average Total Assets
        # Beginning Assets is prior year's ending assets. For the oldest displayed
        # year it comes from the comparative column of the oldest 10-K, which is
        # kept in the statement values even though it is not a displayed year.
        prior_yr = str(int(yr) - 1) if yr.isdigit() else None
        beg_ta = ta_vals.get(prior_yr) if prior_yr else None

        avg_ta = None
        if beg_ta is not None and ta is not None:
            avg_ta = (beg_ta + ta) / 2.0
        elif ta is not None:
            print(f"  [WARNING] FY {yr}: Prior-year (FY {prior_yr}) Total Assets unavailable - "
                  f"Total Asset Turnover shown as N/A.")
        asset_turnover = safe_divide(rev, avg_ta, "Total Asset Turnover", yr)

        # 4. Profitability: Net Profit Margin = Net Income / Revenue
        profit_margin = safe_divide(ni, rev, "Net Profit Margin", yr)

        ratios_by_year[yr] = {
            "current_ratio": current_ratio,
            "debt_to_assets": debt_to_assets,
            "asset_turnover": asset_turnover,
            "profit_margin": profit_margin
        }

        underlying_values_by_year[yr] = {
            "revenue": rev,
            "net_income": ni,
            "current_assets": ca,
            "current_liabilities": cl,
            "total_liabilities": tl,
            "ending_assets": ta,
            "beginning_assets": beg_ta,
            "average_assets": avg_ta,
            "prior_year": prior_yr
        }

    raw_accounts_info = {
        "labels": {
            "revenue": rev_label,
            "net_income": ni_label,
            "current_assets": ca_label,
            "current_liabilities": cl_label,
            "total_liabilities": tl_label,
            "total_assets": ta_label
        },
        "by_year": underlying_values_by_year
    }

    return ratios_by_year, raw_accounts_info


def build_analysis_dataset(years, ratios_by_year, sentiment_results, stock_price_results, available_cats,
                           statement_data=None, filings_metadata=None, requested_count=None):
    """
    Assemble financial ratios, sentiment counts, and stock price data into a unified
    dataset sorted in chronological order (OLDEST to NEWEST) for proper left-to-right charting.

    Every source is keyed by fiscal year (the balance sheet year from SEC
    reportDate), so a 10-K filed in the following calendar year still lines up
    with the statement columns for the year it reports.

    Args:
        years (list): Fiscal years from statement extraction.
        ratios_by_year (dict): Computed financial ratios by year.
        sentiment_results (list): Sentiment metrics from Step 7.
        stock_price_results (list): Historical stock prices from Step 8.
        available_cats (list): Loughran-McDonald sentiment categories.
        statement_data (dict, optional): Step 5 statements, used to check year alignment.
        filings_metadata (list, optional): Filing metadata with 'fiscal_year'.
        requested_count (int, optional): Number of filings the user requested.

    Returns:
        list: List of row dictionaries sorted chronologically (oldest to newest).
    """
    # Sort chronologically (oldest to newest)
    sorted_years = sorted(years, key=lambda y: int(y) if y.isdigit() else y)

    # Index sentiment and stock data by string fiscal year
    sentiment_map = {str(item["year"]): item for item in sentiment_results} if sentiment_results else {}
    stock_map = {str(item["year"]): item for item in stock_price_results} if stock_price_results else {}

    # Year-alignment checks across every worksheet
    year_sets = {}
    for sheet_name in ("Income Statement", "Balance Sheet", "Cash Flows"):
        if statement_data and statement_data.get(sheet_name):
            year_sets[sheet_name] = set(statement_data[sheet_name].get("years", []))
    if sentiment_map:
        year_sets["Sentiment Analysis"] = set(sentiment_map)
    if stock_map:
        year_sets["Stock Prices"] = set(stock_map)
    if filings_metadata:
        year_sets["SEC reportDate fiscal years"] = {
            str(item.get("fiscal_year") or item["year"]) for item in filings_metadata
        }
    analysis_set = set(sorted_years)
    misaligned = False
    for source, source_years in year_sets.items():
        if source_years != analysis_set:
            misaligned = True
            print(f"  [WARNING] Fiscal years in '{source}' ({', '.join(sorted(source_years))}) "
                  f"do not match the Analysis years ({', '.join(sorted(analysis_set))}).")
    if not misaligned:
        print(f"  Fiscal years aligned across all worksheets: {', '.join(sorted_years)}")

    if requested_count and len(sorted_years) < requested_count:
        print(f"  [WARNING] Analysis contains {len(sorted_years)} fiscal year(s), "
              f"fewer than the {requested_count} filing(s) requested.")

    for yr in sorted_years:
        sent = sentiment_map.get(yr)
        if not sent:
            print(f"  [WARNING] FY {yr}: Sentiment data missing.")
        elif sent.get("pn_ratio") is None:
            print(f"  [WARNING] FY {yr}: P/N Ratio unavailable (zero negative words or empty text).")
        if stock_map.get(yr, {}).get("price") is None:
            print(f"  [WARNING] FY {yr}: Stock price missing.")

    # Identify other sentiment categories beyond Positive and Negative
    other_cats = [c for c in available_cats if c not in ("Positive", "Negative")]

    analysis_data = []

    for yr in sorted_years:
        ratios = ratios_by_year.get(yr, {})
        sent = sentiment_map.get(yr, {})
        stock = stock_map.get(yr, {})

        sent_counts = sent.get("counts", {}) if sent else {}

        row = {
            "year": yr,
            "current_ratio": ratios.get("current_ratio"),
            "debt_to_assets": ratios.get("debt_to_assets"),
            "asset_turnover": ratios.get("asset_turnover"),
            "profit_margin": ratios.get("profit_margin"),
            "total_words": sent.get("total_words") if sent else None,
            "positive": sent_counts.get("Positive") if sent else None,
            "negative": sent_counts.get("Negative") if sent else None,
            "other_sentiment": {cat: sent_counts.get(cat) for cat in other_cats},
            "pn_ratio": sent.get("pn_ratio") if sent else None,
            "stock_price": stock.get("price") if stock else None
        }
        analysis_data.append(row)

    return analysis_data


def add_analysis_worksheet_and_chart(company_name, ticker, analysis_data, raw_accounts_info, output_filepath, available_cats):
    """
    Create or safely refresh the 'Analysis' worksheet in the financial analysis Excel
    workbook, write financial ratios, sentiment counts, and stock prices in chronological
    order, document all formulas and metrics, and generate an openpyxl dual-axis
    combination line chart comparing Stock Price ($) vs. MD&A Sentiment P/N Ratio.

    Args:
        company_name (str): Full company name.
        ticker (str): Company ticker symbol.
        analysis_data (list): Chronological list of combined metrics.
        raw_accounts_info (dict): Underpinning account names and values.
        output_filepath (str): Excel workbook file path.
        available_cats (list): Sentiment categories loaded from dictionary.

    Returns:
        bool: True if successful, False otherwise.
    """
    if not os.path.exists(output_filepath):
        print(f"Error: Target workbook '{output_filepath}' not found.")
        return False

    try:
        workbook = openpyxl.load_workbook(output_filepath)
    except Exception as err:
        print(f"Error reading workbook '{output_filepath}': {err}")
        return False

    sheet_title = "Analysis"

    # Safely replace existing Analysis worksheet without touching other sheets
    if sheet_title in workbook.sheetnames:
        del workbook[sheet_title]

    worksheet = workbook.create_sheet(title=sheet_title)

    # Professional styles matching prior tabs
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="595959")
    section_font = Font(name="Calibri", size=11, bold=True, color="1F497D")
    note_font = Font(name="Calibri", size=9, color="333333")
    regular_font = Font(name="Calibri", size=10, color="000000")
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )

    # Determine date range for subtitle
    if analysis_data:
        oldest_yr = analysis_data[0]["year"]
        newest_yr = analysis_data[-1]["year"]
        period_str = f"Fiscal Years {oldest_yr} – {newest_yr} (Chronological Order)"
    else:
        period_str = "Comprehensive Financial & MD&A Sentiment Analysis"

    # Row 1: Title Header
    worksheet["A1"] = f"{company_name.upper()} ({ticker}) — FINANCIAL RATIOS, SENTIMENT & STOCK PRICE ANALYSIS"
    worksheet["A1"].font = title_font

    # Row 2: Subtitle Header
    worksheet["A2"] = f"{period_str} | Unified Valuation & Sentiment Modeling"
    worksheet["A2"].font = subtitle_font

    # Identify other sentiment categories
    other_cats = [c for c in available_cats if c not in ("Positive", "Negative")]

    # Build Header List
    headers = [
        "Fiscal Year",
        "Current Ratio",
        "Debt-to-Assets Ratio",
        "Total Asset Turnover",
        "Net Profit Margin",
        "Total Word Count",
        "Positive Words",
        "Negative Words"
    ]
    for cat in other_cats:
        headers.append(f"{cat} Words")
    headers.extend(["P/N Ratio", "Stock Price"])

    # Row 4: Column Headers
    for col_idx, h in enumerate(headers, 1):
        cell = worksheet.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Dynamic Column Indices
    pn_col_idx = headers.index("P/N Ratio") + 1
    stock_col_idx = headers.index("Stock Price") + 1

    # Rows 5+: Data rows
    for r_offset, row in enumerate(analysis_data):
        r_idx = 5 + r_offset

        # Col 1: Fiscal Year (stored as a number, not text)
        year_val = int(row["year"]) if str(row["year"]).isdigit() else row["year"]
        c = worksheet.cell(row=r_idx, column=1, value=year_val)
        c.font = regular_font
        c.border = cell_border
        c.alignment = Alignment(horizontal="center")

        # Col 2: Current Ratio (Format: 0.00)
        c = worksheet.cell(row=r_idx, column=2, value=row["current_ratio"] if row["current_ratio"] is not None else "N/A")
        c.font = regular_font
        c.border = cell_border
        if row["current_ratio"] is not None:
            c.number_format = "0.00"
            c.alignment = Alignment(horizontal="right")
        else:
            c.alignment = Alignment(horizontal="center")

        # Col 3: Debt-to-Assets Ratio (Format: 0.00%)
        c = worksheet.cell(row=r_idx, column=3, value=row["debt_to_assets"] if row["debt_to_assets"] is not None else "N/A")
        c.font = regular_font
        c.border = cell_border
        if row["debt_to_assets"] is not None:
            c.number_format = "0.00%"
            c.alignment = Alignment(horizontal="right")
        else:
            c.alignment = Alignment(horizontal="center")

        # Col 4: Total Asset Turnover (Format: 0.00)
        c = worksheet.cell(row=r_idx, column=4, value=row["asset_turnover"] if row["asset_turnover"] is not None else "N/A")
        c.font = regular_font
        c.border = cell_border
        if row["asset_turnover"] is not None:
            c.number_format = "0.00"
            c.alignment = Alignment(horizontal="right")
        else:
            c.alignment = Alignment(horizontal="center")

        # Col 5: Net Profit Margin (Format: 0.00%)
        c = worksheet.cell(row=r_idx, column=5, value=row["profit_margin"] if row["profit_margin"] is not None else "N/A")
        c.font = regular_font
        c.border = cell_border
        if row["profit_margin"] is not None:
            c.number_format = "0.00%"
            c.alignment = Alignment(horizontal="right")
        else:
            c.alignment = Alignment(horizontal="center")

        # Col 6: Total Word Count
        c = worksheet.cell(row=r_idx, column=6, value=row["total_words"] if row["total_words"] is not None else "N/A")
        c.font = regular_font
        c.border = cell_border
        if row["total_words"] is not None:
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="right")
        else:
            c.alignment = Alignment(horizontal="center")

        # Col 7: Positive Words
        c = worksheet.cell(row=r_idx, column=7, value=row["positive"] if row["positive"] is not None else "N/A")
        c.font = regular_font
        c.border = cell_border
        if row["positive"] is not None:
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="right")
        else:
            c.alignment = Alignment(horizontal="center")

        # Col 8: Negative Words
        c = worksheet.cell(row=r_idx, column=8, value=row["negative"] if row["negative"] is not None else "N/A")
        c.font = regular_font
        c.border = cell_border
        if row["negative"] is not None:
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="right")
        else:
            c.alignment = Alignment(horizontal="center")

        # Other Sentiment Categories
        current_col = 9
        for cat in other_cats:
            cat_val = row["other_sentiment"].get(cat)
            c = worksheet.cell(row=r_idx, column=current_col, value=cat_val if cat_val is not None else "N/A")
            c.font = regular_font
            c.border = cell_border
            if cat_val is not None:
                c.number_format = "#,##0"
                c.alignment = Alignment(horizontal="right")
            else:
                c.alignment = Alignment(horizontal="center")
            current_col += 1

        # P/N Ratio and Stock Price feed the chart. A missing value is left blank
        # (Excel plots a gap) because a text "N/A" would be plotted as zero.
        c = worksheet.cell(row=r_idx, column=pn_col_idx, value=row["pn_ratio"])
        c.font = regular_font
        c.border = cell_border
        c.number_format = "0.00"
        c.alignment = Alignment(horizontal="right")

        # Stock Price (Format: $#,##0.00)
        c = worksheet.cell(row=r_idx, column=stock_col_idx, value=row["stock_price"])
        c.font = regular_font
        c.border = cell_border
        c.number_format = "$#,##0.00"
        c.alignment = Alignment(horizontal="right")

    # Freeze panes at Row 5 to keep headers pinned
    worksheet.freeze_panes = "A5"

    # Ratio formula documentation to the right of the data table
    notes_col = len(headers) + 2
    worksheet.cell(row=3, column=notes_col, value="RATIO FORMULAS & WHAT THEY MEASURE").font = section_font
    doc_headers = ["Ratio", "Measures", "Formula"]
    doc_rows = [
        ("Current Ratio", "Liquidity", "Current Assets / Current Liabilities"),
        ("Debt-to-Assets", "Solvency", "Total Liabilities / Total Assets"),
        ("Total Asset Turnover", "Efficiency", "Revenue / Average Total Assets"),
        ("Average Total Assets", "(input)", "(Beginning Total Assets + Ending Total Assets) / 2"),
        ("Net Profit Margin", "Profitability", "Net Income / Revenue"),
        ("P/N Ratio", "MD&A Sentiment", "Loughran-McDonald Positive Words / Negative Words"),
        ("Stock Price", "Market Value", "Split-adjusted close on the balance sheet date (or prior trading day)")
    ]
    for offset, h in enumerate(doc_headers):
        cell = worksheet.cell(row=4, column=notes_col + offset, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r_offset, doc_row in enumerate(doc_rows, 5):
        for offset, text in enumerate(doc_row):
            cell = worksheet.cell(row=r_offset, column=notes_col + offset, value=text)
            cell.font = regular_font
            cell.border = cell_border
    note_cell = worksheet.cell(
        row=5 + len(doc_rows) + 1, column=notes_col,
        value="Beginning Total Assets for the oldest year come from the comparative balance sheet in the oldest 10-K. "
              "N/A or blank = data not available (never treated as zero)."
    )
    note_cell.font = note_font
    worksheet.column_dimensions[get_column_letter(notes_col)].width = 22
    worksheet.column_dimensions[get_column_letter(notes_col + 1)].width = 16
    worksheet.column_dimensions[get_column_letter(notes_col + 2)].width = 62

    # Generate Combination Line Chart: Stock Price ($) vs. MD&A Sentiment P/N Ratio
    if len(analysis_data) >= 2:
        chart_start_row = 5 + max(len(analysis_data), len(doc_rows) + 1) + 2

        # Primary Chart: Stock Price ($)
        chart1 = LineChart()
        chart1.title = "Stock Price vs. MD&A Sentiment P/N Ratio"
        chart1.style = 13
        chart1.y_axis.title = "Stock Price ($)"
        chart1.y_axis.number_format = "$#,##0"
        chart1.y_axis.majorGridlines = None
        chart1.x_axis.title = "Fiscal Year"
        chart1.width = 22
        chart1.height = 12
        # Newer Excel versions hide axes unless 'delete' is explicitly False
        chart1.x_axis.delete = False
        chart1.y_axis.delete = False
        chart1.legend.position = "b"

        stock_ref = Reference(worksheet, min_col=stock_col_idx, min_row=4, max_row=4 + len(analysis_data))
        categories = Reference(worksheet, min_col=1, min_row=5, max_row=4 + len(analysis_data))
        chart1.add_data(stock_ref, titles_from_data=True)
        chart1.set_categories(categories)

        # Secondary Chart: Sentiment P/N Ratio
        chart2 = LineChart()
        pn_ref = Reference(worksheet, min_col=pn_col_idx, min_row=4, max_row=4 + len(analysis_data))
        chart2.add_data(pn_ref, titles_from_data=True)
        chart2.y_axis.title = "P/N Ratio"
        chart2.y_axis.axId = 200
        chart2.y_axis.number_format = "0.00"
        chart2.y_axis.crosses = "max"
        chart2.y_axis.delete = False

        # Combine secondary chart into primary chart
        chart1 += chart2

        worksheet.add_chart(chart1, f"B{chart_start_row}")

    # Auto-fit Column Dimensions
    for col_idx, h in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        worksheet.column_dimensions[col_letter].width = max(15, len(h) + 3)

    try:
        workbook.save(output_filepath)
        return True
    except OSError as save_err:
        print(f"\nError saving Excel workbook '{output_filepath}': {save_err}")
        return False


def display_step9_summary(ticker, analysis_data, raw_accounts_info, workbook_path):
    """
    Display a formatted confirmation table for Step 9 in the terminal, including
    both final computed ratios / sentiment / stock prices and the underlying account
    values for accounting student and audit verification.
    """
    print("\n" + "=" * 90)
    print("                  STEP 9: FINAL FINANCIAL & SENTIMENT ANALYSIS")
    print("=" * 90)
    print(f"Company Ticker Symbol : {ticker}")
    print("Analysis Layout       : Chronological (Oldest to Newest)")
    print(f"Excel Workbook Path   : {workbook_path}")
    print("Worksheet Tab Created : Analysis (with Dual-Axis Stock vs. Sentiment Line Chart)")
    print("-" * 90)
    print(f"{'Fiscal Year':<12} {'Current Ratio':<15} {'Debt-to-Assets':<16} {'Asset Turnover':<16} {'Net Margin':<14} {'P/N Ratio':<12} {'Stock Price'}")
    print("-" * 90)

    for row in analysis_data:
        yr = row["year"]
        cr = f"{row['current_ratio']:.2f}" if row["current_ratio"] is not None else "N/A"
        da = f"{row['debt_to_assets'] * 100:.2f}%" if row["debt_to_assets"] is not None else "N/A"
        at = f"{row['asset_turnover']:.2f}" if row["asset_turnover"] is not None else "N/A"
        pm = f"{row['profit_margin'] * 100:.2f}%" if row["profit_margin"] is not None else "N/A"
        pn = f"{row['pn_ratio']:.2f}" if row["pn_ratio"] is not None else "N/A"
        sp = f"${row['stock_price']:,.2f}" if row["stock_price"] is not None else "N/A"

        print(f"{yr:<12} {cr:<15} {da:<16} {at:<16} {pm:<14} {pn:<12} {sp}")

    print("-" * 90)
    print("                UNDERLYING FINANCIAL ACCOUNT VALUES VERIFICATION")
    print("-" * 90)

    # Show exactly which statement line each ratio input came from
    label_names = [
        ("current_assets", "Current Assets"), ("current_liabilities", "Current Liabilities"),
        ("total_liabilities", "Total Liabilities"), ("total_assets", "Total Assets"),
        ("revenue", "Revenue"), ("net_income", "Net Income")
    ]
    for key, display in label_names:
        matched = raw_accounts_info.get("labels", {}).get(key) or "NOT FOUND"
        print(f"  {display:<20} <- '{matched}'")
    print("-" * 90)

    print(f"{'FY':<6} {'Cur Assets':>12} {'Cur Liab':>12} {'Total Liab':>12} {'Total Assets':>13} "
          f"{'Revenue':>12} {'Net Income':>12} {'Prior-Yr TA':>13}")
    print("-" * 90)

    def fmt(value):
        return f"{value:,.0f}" if value is not None else "N/A"

    by_year = raw_accounts_info.get("by_year", {})
    for row in analysis_data:
        yr = row["year"]
        accts = by_year.get(yr, {})
        print(f"{yr:<6} {fmt(accts.get('current_assets')):>12} {fmt(accts.get('current_liabilities')):>12} "
              f"{fmt(accts.get('total_liabilities')):>12} {fmt(accts.get('ending_assets')):>13} "
              f"{fmt(accts.get('revenue')):>12} {fmt(accts.get('net_income')):>12} "
              f"{fmt(accts.get('beginning_assets')):>13}")

    # Note on comparative data for oldest year
    if analysis_data:
        oldest_yr = analysis_data[0]["year"]
        oldest_accts = by_year.get(oldest_yr, {})
        beg_ta = oldest_accts.get("beginning_assets")
        prior_yr = oldest_accts.get("prior_year")
        avg_ta = oldest_accts.get("average_assets")
        if beg_ta is not None and avg_ta is not None and prior_yr:
            print("-" * 90)
            print(f"* Total Asset Turnover for FY {oldest_yr} utilized prior-year FY {prior_yr} ending total assets")
            print(f"  (${beg_ta:,.0f}) from comparative balance sheet data (Average Assets = ${avg_ta:,.0f}).")

    print("=" * 90)
    print("Step 9 complete! Final financial analysis, ratios, and chart saved to Excel workbook.\n")


# ==============================================================================
# PIPELINE ORCHESTRATION & PROGRAMMATIC INTERFACE
# ==============================================================================

def run_analysis_workflow(ticker, filing_count, progress_callback=None, is_cli=False):
    """
    Execute the complete SEC 10-K analysis pipeline (Steps 2 through 9).
    Can be invoked programmatically by Streamlit or executed via the terminal CLI.

    Args:
        ticker (str): Stock ticker symbol (e.g. 'AAPL').
        filing_count (int): Number of filings to retrieve and analyze (1 to 10).
        progress_callback (callable, optional): Callback taking (step_number, stage_name, description).
        is_cli (bool, optional): If True, prints formatted console summary tables.

    Returns:
        dict: Complete structured analysis results or error status.
    """
    ticker = str(ticker).strip().upper()
    try:
        filing_count = int(filing_count)
    except (ValueError, TypeError):
        return {"success": False, "error": f"Invalid filing count '{filing_count}'. Must be an integer between 1 and 10."}

    if not ticker or not re.match(r"^[A-Z0-9\.\-]+$", ticker):
        return {"success": False, "error": f"Invalid ticker symbol '{ticker}'. Please enter a valid stock symbol."}

    if filing_count < 1 or filing_count > 10:
        return {"success": False, "error": f"Filing count ({filing_count}) must be between 1 and 10."}

    # Verify required SEC User-Agent credentials before making any requests to SEC EDGAR
    sec_headers, cred_err = get_sec_credentials()
    if not sec_headers:
        if is_cli:
            print("\n" + "=" * 65)
            print("CONFIGURATION ERROR: SEC Contact Information Required")
            print("=" * 65)
            print(cred_err)
            print("=" * 65 + "\n")
        return {"success": False, "error": cred_err}

    # Helper to report progress safely
    def notify_progress(step_num, stage_name, description):
        if progress_callback:
            try:
                progress_callback(step_num, stage_name, description)
            except Exception:
                pass

    # --------------------------------------------------------------------------
    # Step 1 / 10: Connecting to SEC EDGAR (CIK Lookup)
    # --------------------------------------------------------------------------
    notify_progress(1, "Connecting to SEC EDGAR", f"Locating CIK for ticker '{ticker}' in SEC directory...")
    if is_cli:
        print(f"\nLooking up CIK for ticker '{ticker}' in SEC EDGAR...")

    cik, company_title = get_company_cik(ticker)
    if not cik:
        err_msg = f"Could not find a public company matching ticker '{ticker}' in SEC EDGAR. Please check the ticker symbol."
        if is_cli:
            print(f"\nError: {err_msg}")
        return {"success": False, "error": err_msg}

    if is_cli:
        print(f"Found: {company_title} (CIK: {str(cik).zfill(10)})")

    # --------------------------------------------------------------------------
    # Step 2 / 10: Retrieving 10-K Filings
    # --------------------------------------------------------------------------
    notify_progress(2, "Retrieving 10-K Filings", f"Downloading {filing_count} recent 10-K filings for {company_title}...")
    filings_folder = setup_filing_directory(ticker)
    if not filings_folder:
        return {"success": False, "error": f"Failed to initialize storage folder for ticker '{ticker}'."}

    if is_cli:
        print(f"Retrieving filing history for {company_title}...")
    filings_metadata = get_recent_10k_metadata(cik, filing_count)
    if not filings_metadata:
        err_msg = f"No Form 10-K filings found for {company_title}."
        if is_cli:
            print(err_msg)
        return {"success": False, "error": err_msg}

    download_results = download_10k_filings(ticker, cik, filings_metadata, filings_folder)
    if is_cli:
        display_download_confirmation(company_title, ticker, cik, filings_folder, download_results)

    # --------------------------------------------------------------------------
    # Step 3 / 10: Extracting MD&A (Item 7)
    # --------------------------------------------------------------------------
    notify_progress(3, "Extracting MD&A", "Extracting Item 7 MD&A plain text from filings...")
    if is_cli:
        print("\n" + "-" * 60)
        print("        STARTING STEP 4: MD&A SECTION EXTRACTION")
        print("-" * 60)

    extracted_folder = setup_extracted_text_directory(ticker)
    if not extracted_folder:
        return {"success": False, "error": "Failed to create extracted_text directory."}

    mda_results = process_mda_extractions(ticker, filings_folder, extracted_folder, filings_metadata)
    if is_cli and mda_results:
        display_mda_summary(ticker, extracted_folder, mda_results)

    # --------------------------------------------------------------------------
    # Step 4 / 10: Extracting Financial Statements (Item 8)
    # --------------------------------------------------------------------------
    notify_progress(4, "Extracting Financial Statements", "Extracting Income Statement, Balance Sheet, and Cash Flows...")
    if is_cli:
        print("\n" + "-" * 60)
        print("    STARTING STEP 5: FINANCIAL STATEMENTS EXTRACTION")
        print("-" * 60)

    output_dir = setup_output_directory()
    if not output_dir:
        return {"success": False, "error": "Failed to create output directory."}

    output_filename = f"{ticker}_Financial_Analysis.xlsx"
    output_filepath = os.path.join(output_dir, output_filename)

    statement_data = extract_and_merge_statements(
        filings_folder,
        ticker=ticker,
        filings_metadata=filings_metadata,
        max_periods=filing_count
    )

    if not statement_data:
        err_msg = f"Could not extract financial statements from '{filings_folder}'."
        if is_cli:
            print(err_msg)
        return {"success": False, "error": err_msg}

    success = create_financial_analysis_workbook(company_title, ticker, statement_data, output_filepath)
    if not success:
        return {"success": False, "error": f"Failed to generate initial Excel workbook '{output_filename}'."}

    if is_cli:
        display_step5_summary(ticker, output_filepath, statement_data)

    # --------------------------------------------------------------------------
    # Step 5 / 10: Cleaning MD&A Text
    # --------------------------------------------------------------------------
    notify_progress(5, "Cleaning MD&A Text", "Normalizing Unicode, stripping headers/footers, preserving financial terms...")
    if is_cli:
        print("\n" + "-" * 60)
        print("        STARTING STEP 6: MD&A TEXT CLEANING")
        print("-" * 60)

    cleaned_folder = setup_cleaned_text_directory(ticker)
    if not cleaned_folder:
        return {"success": False, "error": "Failed to initialize cleaned_text directory."}

    clean_results = process_cleaned_mda_files(
        ticker=ticker,
        extracted_folder=extracted_folder,
        cleaned_folder=cleaned_folder,
        filings_metadata=filings_metadata
    )
    if is_cli and clean_results:
        display_step6_summary(ticker, cleaned_folder, clean_results)

    # --------------------------------------------------------------------------
    # Step 6 / 10: Loughran-McDonald Sentiment Analysis
    # --------------------------------------------------------------------------
    notify_progress(6, "Running Loughran-McDonald Sentiment Analysis", "Matching cleaned MD&A tokens against dictionary...")
    if is_cli:
        print("\n" + "-" * 60)
        print("    STARTING STEP 7: LOUGHRAN-MCDONALD SENTIMENT ANALYSIS")
        print("-" * 60)

    dictionaries_folder = setup_dictionaries_directory()
    dict_filepath, dict_filename = find_dictionary_file(dictionaries_folder)
    if not dict_filepath:
        err_msg = f"Loughran-McDonald dictionary file not found in '{dictionaries_folder}'."
        if is_cli:
            display_missing_dictionary_notice(dictionaries_folder)
        return {"success": False, "error": err_msg}

    dict_info = load_loughran_mcdonald_dictionary(dict_filepath)
    if not dict_info or dict_info["total_words"] == 0:
        err_msg = f"Failed to load valid dictionary words from '{dict_filename}'."
        if is_cli:
            print(f"Error: {err_msg}")
        return {"success": False, "error": err_msg}

    sentiment_results = analyze_mda_sentiment(
        cleaned_folder=cleaned_folder,
        ticker=ticker,
        filings_metadata=filings_metadata,
        dict_info=dict_info
    )
    if not sentiment_results:
        err_msg = f"No sentiment analysis results could be generated for ticker '{ticker}'."
        if is_cli:
            print(err_msg)
        return {"success": False, "error": err_msg}

    sheet_success = add_sentiment_worksheet_to_workbook(
        company_name=company_title,
        ticker=ticker,
        sentiment_results=sentiment_results,
        dict_info=dict_info,
        output_filepath=output_filepath
    )

    if is_cli and sheet_success:
        display_step7_summary(
            ticker=ticker,
            dict_info=dict_info,
            workbook_path=output_filepath,
            sentiment_results=sentiment_results
        )

    # --------------------------------------------------------------------------
    # Step 7 / 10: Retrieving Historical Stock Prices
    # --------------------------------------------------------------------------
    notify_progress(7, "Retrieving Historical Stock Prices", "Querying Yahoo Finance for split-adjusted closing prices...")
    if is_cli:
        print("\n" + "-" * 60)
        print("        STARTING STEP 8: STOCK PRICE RETRIEVAL")
        print("-" * 60)

    stock_price_results = retrieve_stock_prices(ticker, filings_metadata)
    if stock_price_results:
        add_stock_prices_worksheet_to_workbook(
            company_name=company_title,
            ticker=ticker,
            stock_results=stock_price_results,
            output_filepath=output_filepath
        )
        if is_cli:
            display_step8_summary(
                ticker=ticker,
                workbook_path=output_filepath,
                stock_results=stock_price_results
            )

    # --------------------------------------------------------------------------
    # Step 8 / 10: Calculating Financial Ratios
    # --------------------------------------------------------------------------
    notify_progress(8, "Calculating Financial Ratios", "Computing Liquidity, Solvency, Efficiency, Profitability...")
    if is_cli:
        print("\n" + "-" * 60)
        print("        STARTING STEP 9: FINAL ANALYSIS & VISUALIZATION")
        print("-" * 60)

    analysis_years = statement_data.get("Income Statement", {}).get("years", [])
    if not analysis_years:
        analysis_years = statement_data.get("Balance Sheet", {}).get("years", [])

    ratios_by_year, raw_accounts_info = calculate_financial_ratios(statement_data, analysis_years)
    available_cats = dict_info.get("categories", [])

    analysis_data = build_analysis_dataset(
        years=analysis_years,
        ratios_by_year=ratios_by_year,
        sentiment_results=sentiment_results,
        stock_price_results=stock_price_results or [],
        available_cats=available_cats,
        statement_data=statement_data,
        filings_metadata=filings_metadata,
        requested_count=filing_count
    )

    # --------------------------------------------------------------------------
    # Step 9 / 10: Creating Excel Workbook (Analysis Tab & Chart)
    # --------------------------------------------------------------------------
    notify_progress(9, "Creating Excel Workbook", f"Adding Analysis tab and dual-axis line chart to '{output_filename}'...")
    analysis_tab_success = add_analysis_worksheet_and_chart(
        company_name=company_title,
        ticker=ticker,
        analysis_data=analysis_data,
        raw_accounts_info=raw_accounts_info,
        output_filepath=output_filepath,
        available_cats=available_cats
    )

    if is_cli and analysis_tab_success:
        display_step9_summary(
            ticker=ticker,
            analysis_data=analysis_data,
            raw_accounts_info=raw_accounts_info,
            workbook_path=output_filepath
        )

    # --------------------------------------------------------------------------
    # Step 10 / 10: Analysis Complete
    # --------------------------------------------------------------------------
    notify_progress(10, "Analysis Complete", "Completed all analysis steps successfully!")

    return {
        "success": True,
        "error": None,
        "ticker": ticker,
        "company_name": company_title,
        "cik": str(cik).zfill(10),
        "filing_count": filing_count,
        "filings_metadata": filings_metadata,
        "statement_data": statement_data,
        "sentiment_results": sentiment_results,
        # retrieve_stock_prices() returns None when Yahoo Finance is unreachable
        "stock_price_results": stock_price_results or [],
        "ratios_by_year": ratios_by_year,
        "raw_accounts_info": raw_accounts_info,
        "analysis_data": analysis_data,
        "available_categories": available_cats,
        "output_filepath": output_filepath,
        "dict_info": dict_info
    }


# ==============================================================================
# MAIN PROGRAM EXECUTION (TERMINAL CLI)
# ==============================================================================

def main():
    """
    Main driver executing Steps 2 through 9 in a unified accounting workflow (CLI).
    """
    print("=" * 65)
    print("   SEC 10-K Retrieval, MD&A Extractor & Financial Analyzer")
    print("                   Steps 2 through 9 Workflow")
    print("=" * 65)
    print("Welcome! Please provide the details below.\n")

    # Step 2: Collect and validate user inputs
    ticker = get_ticker_symbol()
    filing_count = get_filing_count()

    print("\n" + "-" * 40)
    print("          INPUT CONFIRMATION")
    print("-" * 40)
    print(f"Company Ticker Symbol : {ticker}")
    print(f"Number of 10-K Filings: {filing_count}")
    print("-" * 40)

    results = run_analysis_workflow(ticker, filing_count, is_cli=True)
    if not results.get("success"):
        print(f"\nExecution terminated: {results.get('error')}")


if __name__ == "__main__":
    main()
